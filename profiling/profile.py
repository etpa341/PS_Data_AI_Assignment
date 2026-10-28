"""
Kestrel feedback profiling (read-only).

Reads every file and every sheet in ../feedback, profiles structure and fields,
and writes profiling/data_governance.xlsx. No cleaning, no merging, no LLM calls.
Source files are opened read-only and never written.

Run:  python profiling/profile.py
"""
import csv
import hashlib
import io
import re
import statistics
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
FEEDBACK = ROOT / "feedback"
REFERENCE = ROOT / "reference"
OUT = Path(__file__).resolve().parent / "data_governance.xlsx"

WINDOW_START, WINDOW_END = date(2025, 10, 1), date(2026, 9, 30)
MONTHS = pd.period_range("2025-10", "2026-09", freq="M").astype(str).tolist()

# --------------------------------------------------------------------------------------
# Column role dictionary. Order matters: first match wins (comment checked before name).
# --------------------------------------------------------------------------------------
ROLE_RULES = [
    ("comment", r"comment|feedback$|verbatim|customer said|^notes$|^review$|what did they say|customer feedback"),
    ("rating", r"rating|score|csat|stars|how satisfied|how was it"),
    ("date", r"^date|submitted_at|^when$"),
    ("period_month", r"^month$"),
    ("department", r"^dept|department|^area$"),
    ("dealer_site", r"^site$|^dealer$"),
    ("customer_id", r"cust(omer)? no|dms_customer_no"),
    ("record_id", r"^ref$|response_id|booking ref"),
    ("phone", r"phone|mobile"),
    ("email", r"email"),
    ("rego", r"^rego$"),
    ("customer_name", r"^name$|^customer$|reviewer|customer name"),
    ("vehicle", r"vehicle|^model$"),
    ("staff", r"^staff$|taken by|logged by|card given by"),
    ("action", r"action|follow"),
    ("source_channel", r"^source$"),
    ("count_bucket", r"^[1-5]$|grand total"),
]
COLUMN_MEANING = {
    "comment": "Free-text customer feedback (verbatim or as paraphrased by staff)",
    "rating": "Customer satisfaction rating/score",
    "date": "Date of visit / feedback received / survey submitted",
    "period_month": "Reporting month (aggregate table)",
    "department": "Department the feedback relates to",
    "dealer_site": "Dealership / site identifier inside a multi-site file",
    "customer_id": "Internal DMS customer number (pseudonymous identifier)",
    "record_id": "Response / booking / reference ID",
    "phone": "Customer phone number (PII)",
    "email": "Customer email (PII)",
    "rego": "Vehicle registration plate (PII, links to owner)",
    "customer_name": "Customer name (PII)",
    "vehicle": "Vehicle model / brand",
    "staff": "Staff member who took / logged the feedback",
    "action": "Follow-up action taken by the dealership",
    "source_channel": "Collection channel (SMS survey, card, ...)",
    "count_bucket": "Count of responses in a rating bucket (aggregate table)",
    "unknown": "Not mapped - review",
}
STANDARD_FIELD = {  # role -> proposed standard field in the consolidated dataset
    "comment": "comment", "rating": "rating_raw", "date": "feedback_date",
    "department": "department", "dealer_site": "dealer_code", "customer_id": "customer_ref (hashed)",
    "record_id": "source_record_id", "phone": "DROP (PII)", "email": "DROP (PII)",
    "rego": "DROP (PII)", "customer_name": "DROP (PII)", "vehicle": "vehicle_model",
    "staff": "staff_first_name (optional)", "action": "followup_action (optional)",
    "source_channel": "collection_channel", "period_month": "n/a (aggregate)",
    "count_bucket": "n/a (aggregate)", "unknown": "review",
}
PII_ROLES = {"phone", "email", "rego", "customer_name", "customer_id"}


def role_of(header):
    if header is None:
        return None
    h = str(header).strip().lower()
    for role, pat in ROLE_RULES:
        if re.search(pat, h):
            return role
    return "unknown"


def is_header_row(row):
    cells = [v for v in row if v not in (None, "")]
    # header cells are short labels; data rows contain dates, numbers or long free text
    if any(not isinstance(v, str) or len(v) > 40 or date_kind(v)[0] not in ("unparseable_text", "month_label") for v in cells):
        return False
    roles = [role_of(v) for v in cells]
    known = [r for r in roles if r not in (None, "unknown")]
    has_core = any(r in ("date", "rating", "comment", "department", "period_month") for r in known)
    # numeric cells (ratings, serial dates) must not be mistaken for headers
    numeric = sum(isinstance(v, (int, float)) and not isinstance(v, bool) for v in row)
    return len(known) >= 2 and has_core and numeric == 0


# --------------------------------------------------------------------------------------
# Department mapping
# --------------------------------------------------------------------------------------
DEPT_MAP = {
    "sales": "Sales", "new cars": "Sales", "used cars": "Sales", "sales - new": "Sales",
    "sales - used": "Sales", "s": "Sales (suspected)",
    "service": "Service", "svc": "Service", "sv": "Service", "service dept": "Service",
    "workshop": "Service",
    "parts": "Parts", "p": "Parts", "p&a": "Parts", "parts & accessories": "Parts",
    "parts counter": "Parts",
}
SALES_KW = re.compile(r"\b(?:buy|bought|purchas|sales|salesman|salesperson|finance|trade-in|test drive|deposit|delivery|handover|new car|drive-away)", re.I)
SERVICE_KW = re.compile(r"\b(?:servic|mechanic|repair|workshop|logbook|brake|oil|booking|tech|loan car|courtesy)", re.I)
PARTS_KW = re.compile(r"\b(?:part|counter|order|accessor|fitment|globe|filter|back-order)", re.I)


def map_dept(v):
    if v is None:
        return "Unknown"
    return DEPT_MAP.get(str(v).strip().lower(), "Unknown")


# --------------------------------------------------------------------------------------
# Rating handling
# --------------------------------------------------------------------------------------
LABEL_ORDER = {
    "very dissatisfied": 1, "dissatisfied": 2, "neutral": 3, "satisfied": 4, "very satisfied": 5,
    "bad": 1, "ok": 2, "good": 3,
}
STAR = "★"


def rating_value(v):
    """Return (kind, numeric_or_None) without assuming direction."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return "blank", None
    if isinstance(v, bool):
        return "other", None
    if isinstance(v, (int, float)):
        return ("int" if float(v).is_integer() else "float"), float(v)
    s = str(v).strip()
    if STAR in s or "☆" in s:
        return "stars", float(s.count(STAR))
    if re.fullmatch(r"-?\d+(\.\d+)?", s):
        return ("int_text" if "." not in s else "float_text"), float(s)
    if s.lower() in LABEL_ORDER:
        return "label", float(LABEL_ORDER[s.lower()])
    return "other", None


POS_KW = re.compile(r"\b(great|excellent|friendly|helpful|thank|thanks|cheers|fantastic|lovely|happy|recommend|ace|legend|impressed|pleased|smooth|well done|top notch)\b", re.I)
NEG_KW = re.compile(r"\b(rude|never|waited|disappoint\w*|useless|terrible|worst|nobody|no one|overcharg\w*|chas\w+|complain\w*|annoy\w*|unhappy|not happy|ridiculous|appalling|still)\b", re.I)


# --------------------------------------------------------------------------------------
# Dates
# --------------------------------------------------------------------------------------
MON = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def date_kind(v):
    """Classify a raw date cell. Returns (kind, parts) where parts are used for parsing."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return "blank", None
    if isinstance(v, datetime):
        return ("excel_datetime_with_time" if (v.hour or v.minute) else "excel_datetime"), v.date()
    if isinstance(v, date):
        return "excel_date", v
    if isinstance(v, (int, float)) and 30000 < v < 60000:
        return "excel_serial_number", (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
    s = str(v).strip()
    if m := re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[ T]\d{2}:\d{2}(?::\d{2})?)?", s):
        return ("iso_text_with_time" if len(s) > 10 else "iso_text"), date(*map(int, m.groups()))
    if m := re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s):
        return "slash_4digit_year", tuple(map(int, m.groups()))
    if m := re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2})", s):
        a, b, y = map(int, m.groups())
        return "slash_2digit_year", (a, b, 2000 + y)
    if m := re.fullmatch(r"(\d{1,2})[- ]([A-Za-z]{3})[- ](\d{2,4})", s):
        d, mon, y = m.groups()
        y = int(y) + (2000 if len(y) == 2 else 0)
        sep = "-" if "-" in s else " "
        return f"d{sep}Mon{sep}yy", date(y, MON[mon.lower()], int(d))
    if m := re.fullmatch(r"(\d{1,2})-(\d{1,2})-(\d{4})", s):
        return "dash_4digit_year", tuple(map(int, m.groups()))
    if re.fullmatch(r"(early|mid|late|wk\d)\s*[A-Za-z]{3}\??|[A-Za-z]{3}\?", s, re.I):
        return "vague_text (e.g. 'early Oct')", None
    if re.search(r"\bago\b", s, re.I):
        return "relative_text", None
    if re.fullmatch(r"[A-Za-z]{3}-\d{2}", s):
        return "month_label", None
    return "unparseable_text", None


def resolve_dayfirst(kinds_parts):
    """Decide D/M vs M/D for ambiguous slash dates using evidence within the sheet."""
    parts = [p for k, p in kinds_parts if isinstance(p, tuple)]
    first_gt12 = sum(p[0] > 12 for p in parts)
    second_gt12 = sum(p[1] > 12 for p in parts)
    ambiguous = sum(p[0] <= 12 and p[1] <= 12 and p[0] != p[1] for p in parts)
    if first_gt12 and second_gt12:
        order = "MIXED (D/M and M/D both present)"
    elif second_gt12:
        order = "M/D (US) - proven by values with 2nd part > 12"
    elif first_gt12:
        order = "D/M - proven by values with 1st part > 12"
    elif parts:
        order = "UNPROVEN (no part > 12); assumed D/M"
    else:
        order = ""
    return order, first_gt12, second_gt12, ambiguous


def to_date(kind, p, order):
    if p is None:
        return None
    if isinstance(p, date):
        return p
    a, b, y = p
    d, m = (b, a) if order.startswith("M/D") else (a, b)
    try:
        return date(y, m, d)
    except ValueError:
        return None


# --------------------------------------------------------------------------------------
# PII patterns (counts only - never emitted)
# --------------------------------------------------------------------------------------
PHONE_RE = re.compile(r"(?<!\d)(?:\+?61\s?|0)[2-478](?:[ -]?\d){8}(?!\d)")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
REGO_RE = re.compile(r"\brego\s*:?\s*[A-Z0-9]{5,7}\b|\b(?=[A-Z0-9]*\d)(?=[A-Z0-9]*[A-Z])[A-Z0-9]{6}\b")
NAME_PREFIX_RE = re.compile(r"^[A-Z][a-z]+ [A-Z][a-z'\-]+ - ")
MOJIBAKE_RE = re.compile("�|â€|Ã.")


# --------------------------------------------------------------------------------------
# Dealer matching
# --------------------------------------------------------------------------------------
dealers = pd.read_csv(REFERENCE / "dealerships.csv")
revenue = pd.read_csv(REFERENCE / "monthly_revenue.csv")
SUBURB_TO_CODE = {r.suburb.lower().replace(" ", ""): r.dealer_code for r in dealers.itertuples()}
FILE_ALIASES = {  # tokens that are not suburb names - each one is a judgement call
    "hunter": ("NSW05", "SUSPECTED: 'Hunter Motors' ~ Hunter region = Newcastle (NSW05 is the only Newcastle-area site). Confirm via brand mix in Vehicle column vs reference brands and with head office."),
    "ips": ("QLD03", "SUSPECTED: 'IPS' / site code 'KMG-IPS' ~ Ipswich. Confirm with head office."),
    "northqld": ("QLD07|QLD08", "Multi-site file: Site column splits Rockhampton (QLD07) and Townsville (QLD08)."),
}
SITE_ALIASES = {
    "rocky": "QLD07", "rok": "QLD07", "rockhampton": "QLD07",
    "tsv": "QLD08", "townsville": "QLD08",
    "kmg-ips": "QLD03", "hunter": "NSW05",
}


def match_dealer(fname):
    key = re.sub(r"[^a-z]", "", fname.lower())
    hits = [(sub, code) for sub, code in SUBURB_TO_CODE.items() if sub in key]
    if hits:  # longest match wins: 'springwood' also contains 'ringwood'
        sub, code = max(hits, key=lambda h: len(h[0]))
        return code, f"Suburb '{sub}' in file name", False
    for tok, (code, ev) in FILE_ALIASES.items():
        if tok in key:
            return code, ev, True
    return "UNMATCHED", "No suburb or alias found", True


# --------------------------------------------------------------------------------------
# Period from file / sheet name
# --------------------------------------------------------------------------------------
def months_between(a, b):
    return pd.period_range(a, b, freq="M").astype(str).tolist()


def period_from_name(text):
    t = text.lower()
    if m := re.search(r"(20\d\d)[\s_]*q([1-4])|q([1-4])[\s_]*(20\d\d)", t):
        y = int(m.group(1) or m.group(4)); q = int(m.group(2) or m.group(3))
        start = pd.Period(f"{y}-{3*q-2:02d}", "M")
        return months_between(start, start + 2), f"Q{q} {y} (calendar quarter)"
    if m := re.search(r"(20\d\d)-(\d{2})(?!\d)", t):
        p = f"{m.group(1)}-{m.group(2)}"
        return [p], p
    if m := re.search(r"([a-z]{3})(\d{2})-([a-z]{3})(\d{2})", t):
        a = f"20{m.group(2)}-{MON[m.group(1)]:02d}"; b = f"20{m.group(4)}-{MON[m.group(3)]:02d}"
        return months_between(a, b), f"{a} to {b}"
    if m := re.fullmatch(r"([a-z]{3}) (\d{2})", t.strip()):
        p = f"20{m.group(2)}-{MON[m.group(1)]:02d}"
        return [p], p
    if m := re.fullmatch(r"(20\d\d)", t.strip()):
        y = int(m.group(1))
        ms = [x for x in MONTHS if x.startswith(str(y))]
        return ms, f"calendar {y} (within window)"
    if "2025-26" in t or "25-26" in t:
        return MONTHS, "FY 2025-26 (window)"
    if m := re.search(r"(?<!\d)(20\d\d)(?!\d)", t):
        y = int(m.group(1))
        return [x for x in MONTHS if x.startswith(str(y))], f"calendar {y} (within window)"
    return None, ""


# --------------------------------------------------------------------------------------
# Loading
# --------------------------------------------------------------------------------------
def load_csv(path):
    raw = path.read_bytes()
    meta = {"bom": raw.startswith(b"\xef\xbb\xbf"), "line_endings": "CRLF" if b"\r\n" in raw else "LF"}
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(enc)
            meta["encoding"] = "utf-8 (BOM)" if (enc == "utf-8-sig" and meta["bom"]) else ("utf-8" if enc == "utf-8-sig" else "cp1252 (Windows-1252; not valid UTF-8)")
            break
        except UnicodeDecodeError:
            continue
    first = text.splitlines()[0]
    delim = max([",", ";", "\t", "|"], key=first.count)
    meta["delimiter"] = {",": "comma", ";": "semicolon", "\t": "tab", "|": "pipe"}[delim]
    meta["non_ascii_chars"] = sum(ord(c) > 127 for c in text)
    rows = [[(c if c != "" else None) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    meta["field_count_dist"] = dict(Counter(len(r) for r in rows))
    return [("(csv)", rows, meta)]


def load_xlsx(path):
    out = []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    wbf = openpyxl.load_workbook(path, data_only=False, read_only=False)
    for ws in wb.worksheets:
        wsf = wbf[ws.title]
        formulas = sum(1 for row in wsf.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        meta = {
            "sheet_state": ws.sheet_state, "merged_ranges": len(ws.merged_cells.ranges),
            "formula_cells": formulas, "max_row": ws.max_row, "max_col": ws.max_column,
            "encoding": "xlsx (n/a)", "delimiter": "n/a",
        }
        out.append((ws.title, rows, meta))
    return out


def nonempty(row):
    return [v for v in row if v not in (None, "") and not (isinstance(v, str) and not v.strip())]


# --------------------------------------------------------------------------------------
# Profile one sheet
# --------------------------------------------------------------------------------------
def profile_sheet(fname, sheet, rows, meta, dealer_code):
    P = {"file_name": fname, "sheet_name": sheet}
    P.update({k: meta.get(k) for k in ("encoding", "delimiter", "bom", "line_endings", "non_ascii_chars", "sheet_state", "merged_ranges", "formula_cells")})
    P["raw_rows"] = len(rows)

    # segment into blocks by header rows (handles repeated headers / template changes)
    header_idx = [i for i, r in enumerate(rows) if is_header_row(r)]
    # a header-like row directly above another header is a group/banner row (e.g. merged 'Visit | Customer | Feedback')
    group_hdr = [i for i in header_idx if i + 1 in header_idx]
    header_idx = [i for i in header_idx if i not in group_hdr]
    P["header_rows"] = ", ".join(str(i + 1) for i in header_idx) or "NONE FOUND"
    pre = rows[: header_idx[0]] if header_idx else rows
    banners = [" / ".join(map(str, nonempty(r))) if i in group_hdr else str(nonempty(r)[0]) for i, r in enumerate(pre) if nonempty(r)]
    P["banner_rows_text"] = " | ".join(banners)[:500]
    P["banner_row_count"] = len(banners)

    data, blanks, totals, notes, rep_headers = [], [], [], [], []
    templates = []
    for bi, h in enumerate(header_idx):
        end = header_idx[bi + 1] if bi + 1 < len(header_idx) else len(rows)
        header = rows[h]
        templates.append([None if v is None else str(v) for v in header])
        if bi > 0:
            rep_headers.append(h + 1)
        for i in range(h + 1, end):
            r = rows[i]
            ne = nonempty(r)
            if not ne:
                blanks.append(i + 1)
                continue
            first = str(ne[0]).strip().lower()
            if re.match(r"^(grand )?total|^average|^sum\b|^count of", first):
                totals.append(i + 1)
                continue
            dk, _ = date_kind(r[[k for k, v in enumerate(header) if role_of(v) in ("date", "period_month")][0]] if any(role_of(v) in ("date", "period_month") for v in header) else None)
            if len(ne) == 1 and isinstance(ne[0], str) and len(ne[0]) > 15 and dk in ("unparseable_text", "blank"):
                notes.append(i + 1)
                continue
            data.append((i + 1, bi, r))

    test_rows = [rn for rn, _, r in data if any(isinstance(v, str) and re.search(r"\btest(ing)?\b|^asdf$", v.strip(), re.I) and len(v) < 40 for v in r)
                 or any(isinstance(v, str) and v.strip() == "01/01/1900" for v in r)]
    P["test_junk_rows"] = ", ".join(map(str, test_rows))
    P["data_rows"] = len(data)
    P["blank_rows"] = len(blanks)
    P["total_summary_rows"] = ", ".join(map(str, totals))
    P["footer_note_rows"] = ", ".join(map(str, notes))
    P["footer_note_text"] = " | ".join(str(nonempty(rows[n - 1])[0])[:200] for n in notes)
    P["repeated_header_rows"] = ", ".join(map(str, rep_headers))
    P["template_variants"] = len({tuple(t) for t in templates})
    P["columns"] = " || ".join(" | ".join(str(c) for c in t) for t in dict.fromkeys(tuple(t) for t in templates))
    if len({tuple(t) for t in templates}) > 1:
        P["template_change"] = f"YES - header changes at row(s) {P['repeated_header_rows']}"
    elif rep_headers:
        P["template_change"] = f"Repeated identical header at row(s) {P['repeated_header_rows']}"
    else:
        P["template_change"] = "no"

    # columns with data but no header
    unnamed = Counter()
    for _, bi, r in data:
        hdr = templates[bi]
        for k, v in enumerate(r):
            if v not in (None, "") and (k >= len(hdr) or hdr[k] is None):
                unnamed[k + 1] += 1
    P["unnamed_cols_with_data"] = "; ".join(f"col {k}: {n} rows" for k, n in sorted(unnamed.items()))

    # sheet kind
    roles_all = {role_of(c) for t in templates for c in t}
    if "google" in sheet.lower() or "reviewer" in " ".join(str(c).lower() for t in templates for c in t):
        P["sheet_kind"] = "NON-FEEDBACK: copied public reviews"
    elif "count_bucket" in roles_all or "period_month" in roles_all:
        P["sheet_kind"] = "AGGREGATE summary (no row-level records)"
    elif not header_idx and not nonempty(sum(rows, [])):
        P["sheet_kind"] = "EMPTY sheet"
    elif not header_idx:
        P["sheet_kind"] = "NO HEADER FOUND"
    else:
        P["sheet_kind"] = "row-level feedback"

    # gather per-role values (rows keep a block-aware role map)
    def col(role):
        out = []
        for rn, bi, r in data:
            hdr = templates[bi]
            idx = [k for k, c in enumerate(hdr) if role_of(c) == role]
            out.append((rn, r[idx[0]] if idx and idx[0] < len(r) else None))
        return out

    # ---------------- dates
    dates = col("date")
    kinds = [(rn, *date_kind(v)) for rn, v in dates]
    order, f12, s12, amb = resolve_dayfirst([(k, p) for _, k, p in kinds])
    P["date_formats"] = "; ".join(f"{k}: {n}" for k, n in Counter(k for _, k, _ in kinds).most_common())
    P["date_order"] = order
    P["date_ambiguous_dm_count"] = amb
    P["date_part2_gt12"] = s12
    P["date_example_m_d"] = next((f"r{rn}='{v}'" for (rn, v), (_, k, p_) in zip(dates, kinds) if isinstance(p_, tuple) and p_[1] > 12), "")
    parsed = [(rn, to_date(k, p, order)) for rn, k, p in kinds]
    good = [d for _, d in parsed if d]
    P["date_min"] = min(good).isoformat() if good else None
    P["date_max"] = max(good).isoformat() if good else None
    P["date_unparsed"] = sum(1 for (_, k, p), (_, d) in zip(kinds, parsed) if k != "blank" and d is None)
    P["date_blank"] = sum(1 for _, k, _ in kinds if k == "blank")
    exp_months, exp_label = period_from_name(sheet) if sheet != "(csv)" else (None, "")
    if not exp_months:
        exp_months, exp_label = period_from_name(fname)
    P["expected_period_from_name"] = exp_label or "not stated in name -> Oct 2025 to Sep 2026 window"
    exp = set(exp_months or MONTHS)
    oop = [(rn, d) for rn, d in parsed if d and d.strftime("%Y-%m") not in exp]
    oow = [(rn, d) for rn, d in parsed if d and not (WINDOW_START <= d <= WINDOW_END)]
    P["dates_outside_file_period"] = len(oop)
    P["dates_outside_file_period_rows"] = ", ".join(f"r{rn}={d}" for rn, d in oop[:8])
    P["dates_outside_window"] = len(oow)
    P["dates_outside_window_rows"] = ", ".join(f"r{rn}={d}" for rn, d in oow[:8])
    # suspected year typo: an out-of-window date that fits its neighbours once shifted back one year
    seq_all = [d for _, d in parsed]
    typo = 0
    for k, (rn, d) in enumerate(parsed):
        if d and not (WINDOW_START <= d <= WINDOW_END):
            nb = [x for x in seq_all[max(0, k - 3):k] + seq_all[k + 1:k + 4] if x and WINDOW_START <= x <= WINDOW_END]
            try:
                shifted = d.replace(year=d.year - 1)
            except ValueError:
                continue
            if nb and all(abs((shifted - x).days) <= 45 for x in nb):
                typo += 1
    P["dates_suspected_year_typo"] = typo
    seq = [d for _, d in parsed if d]
    P["dates_out_of_sequence"] = sum(1 for a, b in zip(seq, seq[1:]) if b < a)
    months_present = sorted({d.strftime("%Y-%m") for d in good})
    P["months_present"] = ", ".join(months_present)
    P["months_missing_vs_expected"] = ", ".join(m for m in sorted(exp) if m not in months_present) if P["sheet_kind"] == "row-level feedback" else ""

    # ---------------- ratings
    ratings = col("rating")
    rk = [(rn, v, *rating_value(v)) for rn, v in ratings]
    P["rating_kinds"] = "; ".join(f"{k}: {n}" for k, n in Counter(k for _, _, k, _ in rk).most_common())
    P["rating_distinct_values"] = "; ".join(f"{str(v)!s}: {n}" for v, n in Counter(str(v) if v is not None else "<blank>" for _, v, _, _ in rk).most_common(20))
    nums = [x for *_, x in rk if x is not None]
    P["rating_min"] = min(nums) if nums else None
    P["rating_max"] = max(nums) if nums else None
    kinds_set = {k for _, _, k, _ in rk} - {"blank"}
    if "stars" in kinds_set:
        scale = "1-5 star glyphs"
    elif "label" in kinds_set:
        labels = sorted({str(v).strip() for _, v, k, _ in rk if k == "label"})
        scale = "text labels: " + " / ".join(labels)
    elif nums and max(nums) > 5:
        scale = f"numeric {int(min(nums))}-{int(max(nums))} (looks like 0/1-10)"
    elif nums:
        scale = f"numeric {int(min(nums))}-{int(max(nums))} (1-5)"
    else:
        scale = "none"
    if P["template_variants"] > 1 and nums:
        per_block = defaultdict(list)
        for (rn, bi, _), (_, _, _, x) in zip(data, rk):
            if x is not None:
                per_block[bi].append(x)
        scale += " | per template block: " + "; ".join(f"block{b+1} max={max(v):g}" for b, v in per_block.items())
    P["rating_scale_observed"] = scale
    P["rating_stated_on_sheet"] = "; ".join(b for b in banners if re.search(r"rating|=|scale", b, re.I)) + "; ".join(c for t in templates for c in t if c and re.search(r"\(.*\d.*\)", c))
    P["rating_blank"] = sum(1 for _, _, k, _ in rk if k == "blank")
    P["rating_non_parsable"] = sum(1 for _, _, k, _ in rk if k == "other")
    P["rating_float_values"] = sum(1 for _, _, k, _ in rk if k.startswith("float"))

    # ---------------- comments
    comments = col("comment")
    texts = [(rn, str(v)) for rn, v in comments if v not in (None, "") and str(v).strip()]
    P["comment_blank"] = len(comments) - len(texts)
    lens = [len(t) for _, t in texts]
    P["comment_len_min"] = min(lens) if lens else None
    P["comment_len_median"] = statistics.median(lens) if lens else None
    P["comment_len_max"] = max(lens) if lens else None
    P["comment_multiline"] = sum("\n" in t for _, t in texts)
    P["comment_all_caps"] = sum(t.isupper() and len(t) > 15 for _, t in texts)
    P["comment_mojibake"] = sum(bool(MOJIBAKE_RE.search(t)) for _, t in texts)
    P["comment_staff_paraphrase"] = sum(bool(re.match(r"^(cust|customer|said|older|lady|bloke)\b", t, re.I)) for _, t in texts)
    # PII inside free text (counts only)
    P["pii_in_comment_phone"] = sum(bool(PHONE_RE.search(t)) for _, t in texts)
    P["pii_in_comment_email"] = sum(bool(EMAIL_RE.search(t)) for _, t in texts)
    P["pii_in_comment_rego"] = sum(bool(REGO_RE.search(t)) for _, t in texts)
    P["pii_in_comment_name_prefix"] = sum(bool(NAME_PREFIX_RE.match(t)) for _, t in texts)

    # ---------------- inversion check (keyword heuristic, suspected only)
    rmap = {rn: x for rn, _, _, x in rk}
    pos = [rmap[rn] for rn, t in texts if rmap.get(rn) is not None and POS_KW.search(t) and not NEG_KW.search(t)]
    neg = [rmap[rn] for rn, t in texts if rmap.get(rn) is not None and NEG_KW.search(t) and not POS_KW.search(t)]
    P["inversion_pos_kw_mean"] = round(statistics.mean(pos), 2) if pos else None
    P["inversion_neg_kw_mean"] = round(statistics.mean(neg), 2) if neg else None
    P["inversion_n_pos_neg"] = f"{len(pos)}/{len(neg)}"
    if len(pos) >= 8 and len(neg) >= 8:
        gap = statistics.mean(pos) - statistics.mean(neg)
        P["inversion_signal"] = ("SUSPECTED INVERTED (positive-word comments rate lower)" if gap <= -0.5
                                 else "normal direction" if gap >= 0.5 else "inconclusive (gap < 0.5)")
    else:
        P["inversion_signal"] = "insufficient evidence (<8 pos or neg comments)"
    # rating vs comment conflicts (simple, suspected)
    P["rating_comment_conflicts"] = None

    # ---------------- department
    depts = col("department")
    if not any(role_of(c) == "department" for t in templates for c in t):
        dvals = [(rn, sheet) for rn, _ in depts]
        P["department_source"] = "TAB NAME (no department column)"
    else:
        dvals = depts
        P["department_source"] = "column"
    dc = Counter(str(v) if v is not None else "<blank>" for _, v in dvals)
    P["department_raw_values"] = "; ".join(f"{v!r}: {n}" for v, n in dc.most_common())
    P["department_mapped"] = "; ".join(f"{k}: {n}" for k, n in Counter(map_dept(v if v != "<blank>" else None) for v, c in dc.items() for _ in range(c)).most_common())

    # ---------------- dealer / site
    sites = col("dealer_site")
    if any(role_of(c) == "dealer_site" for t in templates for c in t):
        sc = Counter(str(v) for _, v in sites)
        P["site_values"] = "; ".join(f"{v!r}: {n} -> {SITE_ALIASES.get(str(v).strip().lower(), 'UNMAPPED')}" for v, n in sc.most_common())
    else:
        P["site_values"] = ""

    # ---------------- PII columns (counts of populated cells)
    pii = {}
    for role in PII_ROLES:
        if any(role_of(c) == role for t in templates for c in t):
            pii[role] = sum(1 for _, v in col(role) if v not in (None, ""))
    P["pii_columns_populated"] = "; ".join(f"{k}: {v}" for k, v in sorted(pii.items()))

    # ---------------- duplicates within sheet
    keyed = [tuple("" if v is None else str(v).strip() for v in r) for _, _, r in data]
    dupc = Counter(keyed)
    P["exact_duplicate_rows"] = sum(n - 1 for n in dupc.values() if n > 1)
    first_seen = {}
    dup_rows = []
    for (rn, _, _), k in zip(data, keyed):
        if k in first_seen:
            dup_rows.append(f"r{rn}=r{first_seen[k]}")
        else:
            first_seen[k] = rn
    P["exact_duplicate_row_refs"] = ", ".join(dup_rows[:10])
    P["sheet_hash"] = hashlib.md5(repr(sorted(keyed)).encode()).hexdigest()

    # record-level extract for cross-sheet checks (kept in memory only)
    recs = []
    dmap = dict(parsed)
    draw = dict(dates)
    cmap = dict(comments)
    dpmap = dict(dvals)
    smap = dict(sites)
    vmap = dict(col("vehicle"))
    idmap = dict(col("record_id"))
    for rn, bi, r in data:
        recs.append({
            "file": fname, "sheet": sheet, "row": rn, "block": bi, "dealer_file": dealer_code,
            "site": smap.get(rn), "date": dmap.get(rn), "date_raw": draw.get(rn), "rating_raw": next((v for x, v, _, _ in rk if x == rn), None),
            "rating_num": rmap.get(rn), "comment": cmap.get(rn), "dept_raw": dpmap.get(rn),
            "vehicle": vmap.get(rn), "record_id": idmap.get(rn),
        })
    return P, templates, recs


# --------------------------------------------------------------------------------------
# Locked decisions (recorded only - nothing is transformed here)
# Each decision links to Issues_Log via (file regex, issue regex) selectors, resolved to
# the current issue ids at run time, because ids are renumbered on every run.
# Evidence marked {placeholders} is recomputed from the data on every run.
# --------------------------------------------------------------------------------------
DECISIONS = [
    {
        "id": "D1", "status": "locked",
        "decision": "Normalise every rating to 0-100: (x - min) / (max - min) * 100, using the true scale min/max per file/sheet/template block. "
                    "Keep rating_raw and rating_scale on every row; never average raw ratings across scales. Fix the Geelong 2025 inversion (x -> 6 - x) BEFORE normalising. "
                    "Two fixed measures: (1) DISSATISFIED (strict, primary) = normalised score < 40: 1-2 of 5, 1-4 of 10, 'Bad' on Tamworth. "
                    "(2) AT RISK (secondary) = normalised score <= 50: <=3 of 5, <=5 of 10, 'OK' or 'Bad' on Tamworth. "
                    "The mean normalised score is a further secondary measure. "
                    "Both definitions are fixed BEFORE any business analysis and may only be tuned against agreement with AI sentiment labels (V2), never against business results. "
                    "Ringwood uses this same rule for its headline (at risk = <=5 of 10); its <=6 version is a sensitivity row only (D2, V1).",
        "rationale": "Scales differ (1-5, 1-10, 3 labels, stars), so raw means are not comparable. A share of negatives is robust to scale length and easy to explain. "
                     "The midpoint leans negative, so the strict measure alone undercounts unhappy customers; the at-risk measure brackets the answer. Fixing both up front prevents picking whichever definition flatters a result.",
        "evidence": "Midpoint leans negative (keyword lean, neg%/pos%): 3 of 5 = {mid5} (5-point files pooled, excluding inverted/duplicate/non-feedback sheets); "
                    "Tamworth 'OK' = {mid_tam}; Ringwood 5 and 6 of 10 = {mid_ring}. Scales observed: {scale_summary}. "
                    "Ringwood 10-point true range confirmed {ringwood_range} (no 0 values in {ringwood_n} rows; header says 'out of 10'). Geelong 2025: {geelong}.",
        "implementation": "Cleaning step: per (file, sheet, block) assign rating_scale from Scales_and_Labels; map labels (Very dissatisfied..Very satisfied = 1..5; Bad/OK/Good = 1..3; stars = filled-star count); "
                          "reverse Geelong 2025 (6 - x); rating_norm = (x - min)/(max - min)*100 using the stated scale (1-5 / 1-10 / 1-3), not a small file's observed range; "
                          "is_dissatisfied = rating_norm < 40; is_at_risk = rating_norm <= 50. Analysis reports both shares side by side.",
        "risk": "Tamworth has only 3 levels (Bad 0 / OK 50 / Good 100), so its shares are coarser than other scales and 'OK' alone moves a customer into at-risk. "
                "Ringwood headline at-risk = <=5 of 10, the same <=50 rule as every other dealer, so headline numbers stay comparable. "
                "Ringwood's 6 of 10 leans negative ({ring6}), so its headline at-risk share is probably understated; a <=6 version is reported side by side as a sensitivity row only (D2), never as an adopted exception. "
                "Keyword lean is a weak proxy for sentiment; V2 re-tests with AI labels.",
        "issues": [(r"Ringwood", r"10-point scale|Template changes"), (r"Geelong", r"inverted"),
                   (r"Ballarat|Tamworth|bendigo", r"text labels")],
    },
    {
        "id": "D2", "status": "locked (V1 done 2026-10-28: headline kept at 5 or below of 10; 6 or below shown as a sensitivity row only)",
        "decision": "Ringwood 10-point follows D1 for the headline: DISSATISFIED = 1-4, AT RISK = <=5 (same <=50 rule as every dealer). "
                    "A <=6 at-risk version is a SENSITIVITY ROW ONLY, reported side by side with the headline in the analysis; it is not an adopted exception and never replaces the headline. "
                    "Test with AI sentiment labels (V1) how far the headline understates Ringwood.",
        "rationale": "Headline comparability across dealers outweighs a Ringwood-specific fit. Customers often use 5-6 of 10 to express dissatisfaction (NPS treats 0-6 as detractors), and the keyword lean points to 6 -> 7 as the flip, so the sensitivity row shows the likely understatement.",
        "evidence": "Provisional keyword-only lean by Ringwood score (neg%/pos% of comments): {ringwood_lean}. Signal is weak (most comments have no clear keywords) but neg > pos at 5 and 6, and the flip appears at 6 -> 7.",
        "implementation": "Analysis: compute Ringwood at-risk at <=5 (headline) and <=6 (sensitivity) and show both side by side wherever Ringwood appears. "
                          "V1: after the LLM sentiment run, cross-tab Ringwood score x sentiment to confirm 1-4 are mostly negative and to quantify how negative 5 and 6 are; record the result here. The headline rule does not change.",
        "risk": "Ringwood's 6 of 10 leans negative ({ring6}), so the headline at-risk share is probably understated. Ringwood has {ringwood_n} ten-point rows, so per-score n is only 10-50 and the size of the understatement is uncertain.",
        "issues": [(r"Ringwood", r"10-point scale")],
    },
    {
        "id": "D3", "status": "locked",
        "decision": "Bendigo department codes: S -> Sales (inferred), SV -> Service, P -> Parts. Add dept_inferred = true on S rows. Run a sensitivity check excluding S rows (V3) and send Bendigo a data request to confirm the code (R4).",
        "rationale": "SV is already the Service code and S is the only other code; a hand review of the comments supports Sales.",
        "evidence": "Hand review of the {bendigo_s_comments} S comments with text ({bendigo_s_rows} S rows in total): 8 of 19 clearly Sales (car purchase, handover, finance or paint/warranty upsell, 'car yard', features walkthrough), "
                    "0 clearly Service, 1 ambiguous ('car not ready'), the rest generic. This is ONE reviewer's judgement (2026-10-27), not recomputed by the script. Without this mapping Bendigo would have 0 Sales rows.",
        "implementation": "Cleaning step: department = {{'S': 'Sales', 'SV': 'Service', 'P': 'Parts'}} for bendigo feedback.csv; dept_inferred = (raw == 'S'). V3: rerun Sales metrics with and without dept_inferred rows.",
        "risk": "If S actually meant Service, Bendigo Sales would be empty and Bendigo Service understated by {bendigo_s_rows} rows. Single reviewer, no second rater.",
        "issues": [(r"bendigo", r"Department values")],
    },
    {
        "id": "D4", "status": "locked",
        "decision": "Shepparton: move the rows dated Oct-Dec 2026 back one year (to Oct-Dec 2025). Keep date_raw and set date_corrected = true on those rows.",
        "rationale": "Row order only: each of these rows sits in file order among rows dated Oct-Dec 2025. The closure date rules the 2026 dates out but does not say what the right date is, and the footer only states the closure date.",
        "evidence": "{shep_rows}; {shep_typo} of them sit among neighbours dated exactly one year earlier. Footer: '{shep_footer}'. Effect of excluding them instead: {shep_effect}.",
        "implementation": "Cleaning step: for VIC09 rows with date > closed_date, if date - 1 year falls within Oct 2025 - Mar 2026 and matches neighbours, set date = date - 1 year, date_corrected = true; otherwise quarantine.",
        "risk": "Inference from row order. Excluding the 6 rows instead would not materially change results (see evidence), so the choice is low-stakes.",
        "issues": [(r"Shepparton", r"outside the Oct 2025|Closed site|Footer note|Months with no rows")],
    },
    {
        "id": "D5", "status": "locked",
        "decision": "Toowoomba Q4 2025 and Penrith June 2026 are data gaps: report them as gaps and raise data requests (R1, R3). Exclude the Toowoomba Q4 summary from row-level analysis; use it only as a sensitivity check on score counts, never for themes. "
                    "Keep both months in the revenue join so they appear as gaps (feedback = missing), not zeros. Exclude the Penrith 'Jun 26' tab (a copy of 'Jul 26'). "
                    "Separate gap: Toowoomba Sales has NO feedback at all for Q4 2025 (not even in the summary), logged as its own gap and request (R2).",
        "rationale": "Neither month has usable row-level feedback; treating them as zero would bias trends and per-transaction rates. The Sales gap is separate because a raw Q4 export would still not fill it if Sales was never surveyed.",
        "evidence": "Toowoomba Q4 2025 summary: {toow_summary}; departments in summary: {toow_depts}; revenue for QLD06 Oct-Dec 2025 exists ({toow_rev}). Toowoomba 2026 file has only {toow_sales_2026} Sales rows. "
                    "Penrith: {penrith}; revenue for NSW02 2026-06 exists ({pen_rev}).",
        "implementation": "Cleaning step: drop Penrith 'Jun 26' and the Toowoomba summary from the row-level table; build dealer x month x department from monthly_revenue.csv and left-join feedback counts, leaving NaN with gap_reason for QLD06 2025-10..12 (all depts), QLD06 Sales 2025-10..12 ('not surveyed / not supplied'), and NSW02 2026-06.",
        "risk": "Toowoomba's Q4 baseline is missing, so its year trend starts in January. Toowoomba Sales cannot be compared with other dealers (Q4: none; 2026: very few rows).",
        "issues": [(r"Toowoomba Q4", r"Aggregate|Department absent"), (r"Penrith", r"Identical sheets|outside the period|Months with no rows|more than one tab"),
                   (r"\(coverage\)", r"zero feedback rows"), (r"Toowoomba", r"Department almost absent")],
    },
    {
        "id": "D6", "status": "locked",
        "decision": "Essendon: where the same comment, date and score appear in both the Sales and the Service tab, keep the Service copy and exclude the Sales copy "
                    "(excluded = true, exclude_reason recorded). Only rows with comment text are matched; blank-comment rows that share a date and score are not treated as copies.",
        "rationale": "Counting both copies double-counts the feedback. On a hand review most of the texts describe workshop and service visits and none a purchase, so Service is the better department.",
        "evidence": "{ess_cross}. The keyword list is narrow, so the department call rests on a hand review of the 12 texts (one reviewer, 2026-10-27, not recomputed): "
                    "9 clearly describe a service or workshop visit (car fixed, part waits, workshop viewing area, customer lounge, headlight replaced, came back still rattling), 0 a purchase, 3 generic.",
        "implementation": "02_silver.py rule R19, driven by config/file_mappings.yaml (essendon: cross_tab_duplicates keep_sheet Service, drop_sheet Sales). Match key = date_raw + rating_raw + comment_raw, comment non-blank.",
        "risk": "A customer could give the same comment about a purchase and a service on one day; judged unlikely for these texts. If any copy belonged to Sales, Essendon Sales loses at most {ess_n} rows.",
        "issues": [(r"ESSENDON", r"Overlapping records|more than one tab")],
    },
    {
        "id": "D7", "status": "locked",
        "decision": "Ipswich vague dates ('early Oct', 'wk1 Feb', 'Aug?'): set month from the month word, leave feedback_date empty and flag date_vague_month_inferred. The rows stay in month-level trends and theme analysis but not in any day-level analysis.",
        "rationale": "The month word is in the row's own date cell (month comes from row dates, never tab names), and every calendar month occurs exactly once in the Oct 2025 - Sep 2026 window, so the year is unambiguous (Oct-Dec = 2025, Jan-Sep = 2026).",
        "evidence": "{ips_vague}",
        "implementation": "02_silver.py rule R04 (config ipswich: vague_dates infer_month). A neighbour check confirms each inferred month is within one month of the nearest parsed rows in file order.",
        "risk": "'Aug?' carries the logger's own doubt, so the month could be off by one. {ips_n} of {ips_total} Ipswich rows, so any monthly figure moves by a few rows at most.",
        "issues": [(r"IPS_feedback", r"Unparseable date")],
    },
    {
        "id": "D8", "status": "locked (confirm with head office, R5)",
        "decision": "Map 'Hunter Motors feedback 25-26.xlsx' to NSW05 and 'IPS_feedback_log.xlsx' to QLD03. Neither file name contains a suburb, so these are alias mappings; every row carries dq flag dealer_alias.",
        "rationale": "No other dealer fits either file: 'Hunter' is the Hunter region and NSW05 is the only dealer there; 'IPS' / site code 'KMG-IPS' abbreviates Ipswich. Brand mix is consistent with the reference list.",
        "evidence": "{alias_ev}",
        "implementation": "config/file_mappings.yaml: hunter dealer_code NSW05, ipswich dealer_code QLD03, both dealer_alias: true; 02_silver.py rule R02 flags the rows.",
        "risk": "If either mapping is wrong, that file's results belong to another site and NSW05 / QLD03 would have no feedback at all. {alias_rows} rows depend on it; head office to confirm (R5).",
        "issues": [(r"Hunter|IPS_feedback", r"Dealer mapped by alias")],
    },
    {
        "id": "D9", "status": "locked",
        "decision": "Q3 headline revenue at risk = Service + Parts only (recurring revenue). Sales is reported on a separate line labelled "
                    "'indicative future repeat and referral sales' and never added to the headline. "
                    "Headline: annual revenue (dealer, department) x dissatisfied share x 5% (low) / 10% (base); at-risk share x 20% (high). "
                    "Sales line: annual sales revenue x dissatisfied share x 1% (low) / 2.5% (base); at-risk share x 5% (high). All rates are labelled assumptions.",
        "rationale": "A dissatisfied service or parts customer can take the next job elsewhere, so recurring revenue is what is at risk. "
                     "A car purchase is already made; what is at risk is the next purchase years away and referrals, which is much less certain. "
                     "Applying a leave rate to the whole purchase price lets Sales dominate the total.",
        "evidence": "{q3_naive}",
        "implementation": "analysis.ipynb Q3: headline table Service + Parts; Sales on its own line with its own rate range; sensitivity tables for both. "
                          "Shares from included silver rows; dealer-department cells with fewer than 30 rated responses use the group department share.",
        "risk": "Both rate ranges are assumptions with no churn or repurchase data behind them. Leaving Sales out of the headline understates value if dissatisfied buyers also stop servicing with Kestrel; "
                "that overlap is already partly counted in the Service line.",
        "issues": [(r"monthly_revenue", r"Revenue concentrated")],
    },
]
VERIFICATION_PLAN = [
    {"id": "V1", "decision": "D2", "step": "Ringwood headline (<=5) vs sensitivity (<=6)",
     "method": "After LLM sentiment labelling, cross-tab Ringwood score (1-10) x sentiment label. Confirm 1-4 are mostly negative; measure the negative-label share at 5 and 6 to size the headline understatement.",
     "pass_criteria": "Headline stays <=5 (D1 rule); Ringwood at-risk reported at <=5 and <=6 side by side, with the gap stated.",
     "status": "done 2026-10-28: 62 of 63 Ringwood scores of 1 to 4 are AI negative or mixed; headline stays at 5 or below, at risk 25.5% vs 32.1% at 6 or below (outputs/d2_ringwood_sensitivity.csv, outputs/evaluation/ringwood_d2.csv)"},
    {"id": "V2", "decision": "D1", "step": "Strict and at-risk definition check",
     "method": "Using the same labels, measure the negative-label share at 3 of 5, 'OK' and 5-6 of 10. Confirm the strict (<40) and at-risk (<=50) definitions agree with sentiment at least as well as the alternatives.",
     "pass_criteria": "Definitions fixed before business analysis; dealer and department rankings compared under strict vs at-risk, and any rank changes called out.",
     "status": "done 2026-10-28: each D1 cutoff is best on its own target (F1 0.87 for strict vs AI negative, 0.94 for at risk vs negative or mixed; best alternative 0.71 and 0.84; outputs/evaluation/v2_cutoff_agreement.csv). A 3 of 5 is never positive (35% negative, 65% mixed; v2_sentiment_by_score.csv). Department order is the same; dealer ranks correlate 0.93, the top 5 differs by one dealer (Chatswood 5th strict, 13th at risk; Chermside enters) and 8 dealers move 3 or more places (v2_rank_comparison.csv)"},
    {"id": "V3", "decision": "D3", "step": "Bendigo S-code sensitivity",
     "method": "Compute Bendigo and group Sales metrics with and without rows where dept_inferred = true.",
     "pass_criteria": "Report whether any Sales conclusion changes when S rows are excluded.",
     "status": "done 2026-10-28: group Sales dissatisfied 10.1% (n=1,527) with Bendigo S rows, 10.2% (n=1,495) without; no Sales conclusion changes (outputs/q1_bendigo_s_sensitivity.csv)"},
]
DATA_REQUESTS = [
    {"id": "R1", "decision": "D5", "dealer": "QLD06 Toowoomba", "request": "Row-level feedback export for Oct-Dec 2025 (date, department, rating, comment); only an aggregate summary was supplied.", "status": "to send"},
    {"id": "R2", "decision": "D5", "dealer": "QLD06 Toowoomba", "request": "Was Sales feedback collected in Oct-Dec 2025? The Q4 summary has no Sales rows and 2026 has very few. If collected, please send it; if not, confirm Sales is not surveyed.", "status": "to send"},
    {"id": "R3", "decision": "D5", "dealer": "NSW02 Penrith", "request": "June 2026 feedback: the 'Jun 26' tab is a copy of 'Jul 26'. Please resend June 2026.", "status": "to send"},
    {"id": "R5", "decision": "D8", "dealer": "NSW05 / QLD03", "request": "Please confirm that 'Hunter Motors feedback 25-26.xlsx' is NSW05 and 'IPS_feedback_log.xlsx' (site code KMG-IPS) is QLD03.", "status": "to send"},
    {"id": "R4", "decision": "D3", "dealer": "VIC08 Bendigo", "request": "Please confirm what department code 'S' means (we assume Sales; 'SV' = Service, 'P' = Parts).", "status": "to send"},
]

POS_NEG = lambda t: "neg" if NEG_KW.search(t) and not POS_KW.search(t) else ("pos" if POS_KW.search(t) and not NEG_KW.search(t) else "none")


def lean_by_score(df):
    """Keyword-only neg%/pos% of comments per rating value (a provisional signal, not sentiment)."""
    df = df[df.comment.notna() & df.rating_num.notna()]
    lean = df.comment.astype(str).map(POS_NEG)
    tab = pd.crosstab(df.rating_num.astype(int), lean)
    return "; ".join(f"{k}: {100*r.get('neg', 0)/r.sum():.0f}/{100*r.get('pos', 0)/r.sum():.0f} (n={r.sum()})" for k, r in tab.iterrows())


def build_decisions(rl, prof_df, iss_df, dept_ev, scales_df, brand_check):
    # rating scale per template block
    blk = rl.groupby(["file", "sheet", "block"]).rating_num.agg(["min", "max", "size"]).reset_index()
    ring10 = blk[blk.file.str.startswith("Ringwood") & (blk["max"] > 5)]
    ring_rows = rl.merge(ring10[["file", "sheet", "block"]], on=["file", "sheet", "block"])
    geel = prof_df[(prof_df.file_name.str.startswith("Geelong")) & (prof_df.sheet_name == "2025")].iloc[0]
    # 5-point pool for the midpoint check: excludes inverted, 10-point, 3-label, duplicate and non-feedback sheets
    excl = (rl.file.str.startswith(("Geelong", "Tamworth", "Chermside Q1 2026 FINAL.xlsx"))
            | ((rl.file.str.startswith("Penrith")) & (rl.sheet == "Jun 26")))
    five = rl[~excl].merge(blk[blk["max"] <= 5][["file", "sheet", "block"]], on=["file", "sheet", "block"])
    tam = rl[rl.file.str.startswith("Tamworth")]
    mid = f"5-point pooled score {lean_by_score(five[five.rating_num == 3])}; Tamworth OK (2 of 3) {lean_by_score(tam[tam.rating_num == 2])}; Ringwood score {lean_by_score(ring_rows[ring_rows.rating_num.isin([5, 6])])} (neg%/pos%)"
    ev_s = dept_ev.get("S", {})
    n_s = ev_s.get("n_comments", 0)
    shep = prof_df[prof_df.file_name.str.startswith("Shepparton")].iloc[0]
    toow = prof_df[prof_df.file_name.str.startswith("Toowoomba Q4")].iloc[0]
    rev_q = revenue[(revenue.dealer_code == "QLD06") & revenue.month.isin(["2025-10", "2025-11", "2025-12"])]
    rev_p = revenue[(revenue.dealer_code == "NSW02") & (revenue.month == "2026-06")]
    pen_issue = iss_df[iss_df.file.str.contains("Penrith") & iss_df.issue.str.contains("Identical")]
    # D4: what happens to Shepparton shares if the 6 post-closure rows are excluded instead of corrected
    sh = rl[rl.dealer_code == "VIC09"].dropna(subset=["rating_num"])
    late = sh.date.astype(str) > str(dealers.set_index("dealer_code").loc["VIC09", "closed_date"])
    def shares(d):
        return f"dissatisfied {100*(d.rating_num <= 2).mean():.1f}%, at-risk {100*(d.rating_num <= 3).mean():.1f}%, mean {d.rating_num.mean():.2f} (n={len(d)})"
    shep_effect = (f"all Shepparton rows {shares(sh)} vs excluding the {int(late.sum())} rows {shares(sh[~late])}; "
                   f"they are {int(late.sum())} of {len(rl)} group rows ({100*late.sum()/len(rl):.2f}%)")
    toow26 = rl[rl.file.str.startswith("Toowoomba feedback 2026")]
    # D6: Essendon Sales rows whose (date, score, comment) also appear in the Service tab
    ess = rl[rl.file.str.startswith("ESSENDON") & rl.ckey.notna()]
    svc_keys = set(zip(ess[ess.sheet == "Service"].date.astype(str), ess[ess.sheet == "Service"].rating_raw.astype(str), ess[ess.sheet == "Service"].ckey))
    sales = ess[ess.sheet == "Sales"]
    cross = sales[[x in svc_keys for x in zip(sales.date.astype(str), sales.rating_raw.astype(str), sales.ckey)]]
    # D7: Ipswich vague date cells
    ips = rl[rl.file.str.startswith("IPS_feedback")]
    vague = ips[ips.date.isna() & ips.date_raw.notna()]
    # D8: alias-mapped files, with dealership names from dealerships.csv
    names = dealers.set_index("dealer_code")
    def alias_line(prefix, code):
        d = rl[rl.file.str.startswith(prefix)]
        b = next((x for x in brand_check if x["dealer_code"] == code), None)
        bt = (f"brands seen {b['brands_seen']} vs reference {b['reference']}" + (f", unexpected {b['unexpected']}" if b["unexpected"] else ", consistent")) if b else "no vehicle column, brand check not possible"
        return (f"{d.file.iloc[0]}: {len(d)} rows, site values {dict(Counter(d.site.astype(str)))} -> {code} '{names.loc[code, 'dealership_name']}' "
                f"({names.loc[code, 'suburb']}, {names.loc[code, 'state']}, region {names.loc[code, 'region']}); {bt}")
    # D9: naive Q3 base case (all departments x dissatisfied share x 10%) from silver, as in analysis.ipynb
    silver_p = ROOT / "data" / "silver" / "feedback_silver.csv"
    if silver_p.exists():
        sv = pd.read_csv(silver_p, dtype=str, keep_default_na=False)
        sv = sv[(sv.excluded == "False") & (sv.score_norm != "")]
        sv["dis"] = sv.dissatisfied == "True"
        cell = sv.groupby(["dealer_code", "department"]).dis.agg(["mean", "size"])
        grp = sv.groupby("department").dis.mean()
        cols = {"Sales": "sales_revenue_aud", "Service": "service_revenue_aud", "Parts": "parts_revenue_aud"}
        base = {}
        for d, c in cols.items():
            annual = revenue.groupby("dealer_code")[c].sum()
            sh = [cell.loc[(k, d), "mean"] if (k, d) in cell.index and cell.loc[(k, d), "size"] >= 30 else grp[d] for k in annual.index]
            base[d] = float((annual.values * sh * 0.10).sum())
        tot = sum(base.values())
        rev_tot = {d: revenue[c].sum() for d, c in cols.items()}
        q3_naive = (f"Naive base case (every department x dissatisfied share x 10%) = ${tot/1e6:.2f}m, of which Sales ${base['Sales']/1e6:.2f}m "
                    f"({100*base['Sales']/tot:.0f}%), Service ${base['Service']/1e6:.2f}m, Parts ${base['Parts']/1e6:.2f}m. "
                    f"Sales is {100*rev_tot['Sales']/sum(rev_tot.values()):.0f}% of group revenue but its dissatisfied share is the lowest of the three "
                    f"({100*grp['Sales']:.1f}% vs Service {100*grp['Service']:.1f}%, Parts {100*grp['Parts']:.1f}%), so the total is driven by purchase prices, not by dissatisfaction")
    else:
        q3_naive = "MISSING (run 02_silver.py first)"
    fill = {
        "mid5": lean_by_score(five[five.rating_num == 3]).split(": ", 1)[1],
        "mid_tam": lean_by_score(tam[tam.rating_num == 2]).split(": ", 1)[1],
        "mid_ring": lean_by_score(ring_rows[ring_rows.rating_num.isin([5, 6])]),
        "ring6": "keyword lean neg/pos " + lean_by_score(ring_rows[ring_rows.rating_num == 6]).split(": ", 1)[1],
        "bendigo_s_rows": int((rl.file.str.startswith("bendigo") & (rl.dept_raw == "S")).sum()),
        "shep_effect": shep_effect,
        "toow_depts": toow.department_raw_values,
        "toow_sales_2026": int((toow26.dept_raw.map(map_dept) == "Sales").sum()),
        "scale_summary": "; ".join(f"{v} ({n} rows)" for v, n in zip(scales_df[scales_df.type == "RATING SCALE"].value, scales_df[scales_df.type == "RATING SCALE"]["count"])),
        "ringwood_range": f"{ring_rows.rating_num.min():g}-{ring_rows.rating_num.max():g}",
        "ringwood_n": len(ring_rows),
        "geelong": f"banner '{geel.banner_rows_text.split('|')[-1].strip()}'; positive-keyword comments average {geel.inversion_pos_kw_mean} vs negative {geel.inversion_neg_kw_mean}",
        "midpoint": mid,
        "ringwood_lean": lean_by_score(ring_rows),
        "bendigo_kw": f"Keyword evidence on {n_s} S comments with text: sales words {ev_s.get('sales_kw_%')}% ({round(n_s*ev_s.get('sales_kw_%', 0)/100)}), service words {ev_s.get('service_kw_%')}% ({round(n_s*ev_s.get('service_kw_%', 0)/100)}), parts words {ev_s.get('parts_kw_%')}%",
        "bendigo_s_comments": n_s,
        "shep_rows": f"{shep.dates_outside_window} rows dated after the window/closure: {shep.dates_outside_window_rows}",
        "shep_typo": shep.dates_suspected_year_typo,
        "shep_footer": shep.footer_note_text,
        "toow_summary": f"{toow.data_rows} summary rows, header '{toow.columns}', no comments",
        "toow_rev": f"service jobs {int(rev_q.service_jobs.sum())}, vehicles {int(rev_q.vehicles_sold.sum())}",
        "penrith": pen_issue.evidence.iloc[0] if len(pen_issue) else "MISSING",
        "pen_rev": f"service jobs {int(rev_p.service_jobs.sum())}, vehicles {int(rev_p.vehicles_sold.sum())}",
        "ess_cross": (f"{len(cross)} Sales-tab rows have the same date, score and comment as a Service-tab row (Sales {', '.join(f'r{r}' for r in cross.row)}); "
                      f"{int(cross.ckey.str.contains(SERVICE_KW).sum())} of them contain service keywords, {int(cross.ckey.str.contains(SALES_KW).sum())} contain sales keywords"),
        "ess_n": len(cross),
        "ips_vague": (f"{len(vague)} of {len(ips)} Ipswich rows have a vague date: {', '.join(repr(str(v)) for v in vague.date_raw)} "
                      f"(rows {', '.join(f'r{r}' for r in vague.row)}); every other Ipswich date parses as D/M"),
        "ips_n": len(vague), "ips_total": len(ips),
        "alias_ev": alias_line("Hunter", "NSW05") + " || " + alias_line("IPS_feedback", "QLD03"),
        "alias_rows": int(rl.file.str.startswith(("Hunter", "IPS_feedback")).sum()),
        "q3_naive": q3_naive,
    }
    rows = []
    for d in DECISIONS:
        ids = []
        for frx, irx in d["issues"]:
            hit = iss_df[iss_df.file.str.contains(frx, regex=True) & iss_df.issue.str.contains(irx, regex=True)]
            ids += [i for i in hit.id if i not in ids]
        files = sorted({f.strip() for f in iss_df[iss_df.id.isin(ids)].file for f in f.split("|")} - {"(coverage)"})
        rows.append({"id": d["id"], "decision": d["decision"], "rationale": d["rationale"],
                     "evidence": d["evidence"].format(**fill), "files_affected": "; ".join(files),
                     "implementation": d["implementation"].format(**fill), "status": d["status"], "limit_or_risk": d["risk"].format(**fill),
                     "linked_issue_ids": ", ".join(ids), "n_selectors": len(d["issues"]),
                     "selectors_matched": sum(bool((iss_df.file.str.contains(f) & iss_df.issue.str.contains(i)).any()) for f, i in d["issues"])})
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------------------
def main():
    files = sorted(p for p in FEEDBACK.iterdir() if p.is_file())
    manifest, profiles, coldict, all_recs, errors, lockfiles = [], [], [], [], [], []
    alias_files = {}
    for path in files:
        code, evidence, needs = match_dealer(path.name)
        if path.name.startswith("~$"):  # Excel owner/lock file: workbook is open in Excel, not data
            lockfiles.append(path.name)
            manifest.append({"file_name": path.name, "sheet_name": "", "file_type": path.suffix.lstrip("."), "dealer_code": "n/a",
                             "row_count": 0, "sheet_kind": "NOT DATA: Excel lock file",
                             "notes": f"{path.stat().st_size}-byte Excel owner file created while the workbook is open; not part of the 52-file pack"})
            continue
        try:
            sheets = load_csv(path) if path.suffix.lower() == ".csv" else load_xlsx(path)
        except Exception as e:  # record, never skip silently
            errors.append((path.name, repr(e)))
            manifest.append({"file_name": path.name, "sheet_name": "", "file_type": path.suffix, "dealer_code": code, "notes": f"UNREADABLE: {e!r}"})
            continue
        for sheet, rows, meta in sheets:
            try:
                P, templates, recs = profile_sheet(path.name, sheet, rows, meta, code)
            except Exception as e:
                errors.append((f"{path.name}::{sheet}", repr(e)))
                manifest.append({"file_name": path.name, "sheet_name": sheet, "file_type": path.suffix, "dealer_code": code, "notes": f"PROFILE ERROR: {e!r}"})
                continue
            P["file_type"] = path.suffix.lstrip(".")
            P["dealer_code"] = code
            P["field_count_dist"] = str(meta.get("field_count_dist", ""))
            profiles.append(P)
            all_recs.extend(recs)
            for t in dict.fromkeys(tuple(t) for t in templates):
                for pos, c in enumerate(t, 1):
                    r = role_of(c) if c is not None else None
                    coldict.append({
                        "file_name": path.name, "sheet_name": sheet, "col_position": pos,
                        "raw_column": c if c is not None else "<no header>",
                        "role": r or "unnamed", "plain_english_meaning": COLUMN_MEANING.get(r or "unknown", "No header - see unnamed_cols_with_data"),
                        "proposed_standard_field": STANDARD_FIELD.get(r or "unknown", "review"),
                        "contains_pii": "YES" if r in PII_ROLES else ("POSSIBLE (free text)" if r == "comment" else "no"),
                    })
            notes = []
            if needs and "|" not in code:
                alias_files[path.name] = (code, evidence)
            if needs:
                notes.append("NEEDS DECISION: " + evidence)
            if P["site_values"]:
                notes.append("Multi-site via Site column")
            if P["sheet_kind"] != "row-level feedback":
                notes.append(P["sheet_kind"])
            manifest.append({
                "file_name": path.name, "sheet_name": sheet, "file_type": P["file_type"],
                "dealer_code": code + (" (needs decision)" if needs else ""),
                "dealer_match_evidence": evidence,
                "period_covered": f"{P['date_min']} to {P['date_max']}" if P["date_min"] else "n/a",
                "expected_period_from_name": P["expected_period_from_name"],
                "row_count": P["data_rows"], "sheet_kind": P["sheet_kind"],
                "notes": "; ".join(notes),
            })

    rec = pd.DataFrame(all_recs)
    rec["dealer_code"] = rec.apply(lambda r: SITE_ALIASES.get(str(r.site).strip().lower(), r.dealer_file) if r.site is not None else r.dealer_file, axis=1)
    rec["month"] = rec["date"].map(lambda d: d.strftime("%Y-%m") if d else None)
    sheet_kind = {(p["file_name"], p["sheet_name"]): p["sheet_kind"] for p in profiles}
    rec["sheet_kind"] = [sheet_kind[(f, s)] for f, s in zip(rec.file, rec.sheet)]
    rl = rec[rec.sheet_kind == "row-level feedback"].copy()

    # ---------------- cross-sheet duplicate / overlap analysis
    rl["ckey"] = rl.comment.map(lambda c: re.sub(r"\s+", " ", str(c).strip().lower()) if c not in (None, "") else None)
    sheets_ = rl.groupby(["file", "sheet"])
    keyset = {k: set(zip(g.date.astype(str), g.rating_raw.astype(str), g.ckey.fillna(""))) for k, g in sheets_}
    hashes = {(p["file_name"], p["sheet_name"]): p["sheet_hash"] for p in profiles}
    overlaps = []
    ks = list(keyset)
    dealer_of = rl.groupby(["file", "sheet"]).dealer_file.first().to_dict()
    for i in range(len(ks)):
        for j in range(i + 1, len(ks)):
            a, b = ks[i], ks[j]
            if dealer_of[a] != dealer_of[b]:
                continue
            inter = keyset[a] & keyset[b]
            if inter:
                overlaps.append({"a": a, "b": b, "shared": len(inter), "a_rows": len(keyset[a]), "b_rows": len(keyset[b]),
                                 "identical": hashes[a] == hashes[b]})
    # same comment text + date across tabs ignoring rating/dept (e.g. escalations copied between tabs)
    cross_tab = []
    for f, g in rl[rl.ckey.notna()].groupby("file"):
        if g.sheet.nunique() < 2:
            continue
        k = g.groupby(["date", "ckey"]).sheet.nunique()
        n = int((k > 1).sum())
        if n:
            ex = g.set_index(["date", "ckey"]).loc[k[k > 1].index[0]]
            cross_tab.append((f, n, "; ".join(f"{s} r{r}" for s, r in zip(ex.sheet, ex.row))))
    # comment texts reused across different dealers
    long_c = rl[rl.ckey.str.len() >= 40] if rl.ckey.notna().any() else rl.iloc[0:0]
    reuse = long_c.groupby("ckey").dealer_code.nunique()
    cross_dealer_reuse = int((reuse > 1).sum())
    # record ID collisions across files
    ids = rl[rl.record_id.notna()].copy()
    ids["record_id"] = ids.record_id.astype(str)
    id_coll = ids.groupby("record_id").file.nunique()
    id_coll_files = ids[ids.record_id.isin(id_coll[id_coll > 1].index)].groupby("file").size().to_dict()
    id_dupe_within = ids.groupby(["file", "record_id"]).size()
    id_dupe_within = id_dupe_within[id_dupe_within > 1].groupby(level=0).size().to_dict()

    # ---------------- brand evidence per dealer (validates dealer mapping)
    brand_prefix = {}
    for v in rl.vehicle.dropna().astype(str):
        parts = v.split()
        if len(parts) >= 2 and parts[0] in ("Altura", "Kodiak", "Veyra", "Voltra"):
            brand_prefix[" ".join(parts[1:]).lower()] = parts[0]
    def brand(v):
        if v is None: return None
        s = str(v)
        p = s.split()
        if p and p[0] in ("Altura", "Kodiak", "Veyra", "Voltra"): return p[0]
        return brand_prefix.get(s.lower())
    rl["brand"] = rl.vehicle.map(brand)
    ref_brands = dealers.set_index("dealer_code").brands.str.split("|").map(set).to_dict()
    brand_check = []
    for code, g in rl[rl.brand.notna()].groupby("dealer_code"):
        seen = Counter(g.brand)
        unexpected = {b: n for b, n in seen.items() if b not in ref_brands.get(code, set())}
        brand_check.append({"dealer_code": code, "brands_seen": dict(seen), "reference": "|".join(sorted(ref_brands.get(code, []))), "unexpected": unexpected})

    # ---------------- department keyword evidence for ambiguous codes
    dept_ev = []
    for raw, g in rl[rl.ckey.notna()].groupby(rl.dept_raw.astype(str)):
        n = len(g)
        dept_ev.append({"raw": raw, "n_comments": n,
                        "sales_kw_%": round(100 * g.ckey.str.contains(SALES_KW).mean(), 1),
                        "service_kw_%": round(100 * g.ckey.str.contains(SERVICE_KW).mean(), 1),
                        "parts_kw_%": round(100 * g.ckey.str.contains(PARTS_KW).mean(), 1)})
    dept_ev = {d["raw"]: d for d in dept_ev}

    # ---------------- coverage (dealer x month)
    cov = rl[rl.month.notna()].groupby(["dealer_code", "month"]).size().unstack(fill_value=0)
    cov = cov.reindex(index=dealers.dealer_code, columns=MONTHS, fill_value=0)
    rev_months = revenue.groupby("dealer_code").month.apply(set).to_dict()

    # ---------------- assemble workbook frames
    prof_df = pd.DataFrame(profiles)
    lead = ["file_name", "sheet_name", "file_type", "dealer_code", "sheet_kind", "raw_rows", "data_rows", "blank_rows",
            "header_rows", "columns", "banner_row_count", "banner_rows_text", "total_summary_rows", "repeated_header_rows",
            "template_change", "template_variants", "footer_note_rows", "footer_note_text", "unnamed_cols_with_data",
            "encoding", "delimiter", "bom", "line_endings", "non_ascii_chars", "field_count_dist", "sheet_state", "merged_ranges", "formula_cells"]
    prof_df = prof_df[lead + [c for c in prof_df.columns if c not in lead and c != "sheet_hash"]]

    man_df = pd.DataFrame(manifest)

    # Scales and labels tab
    scale_rows = []
    for (sc), g in prof_df[prof_df.data_rows > 0].groupby("rating_scale_observed"):
        scale_rows.append({"type": "RATING SCALE", "value": sc, "maps_to": "", "count": int(g.data_rows.sum()),
                           "files": "; ".join(f"{f} [{s}]" for f, s in zip(g.file_name, g.sheet_name)),
                           "evidence": "; ".join(sorted({x for x in g.inversion_signal if x})), "decision_needed": ""})
    dept_counts = Counter()
    dept_files = defaultdict(set)
    for r in rl.itertuples():
        raw = r.dept_raw if r.dept_raw is not None else "<blank>"
        dept_counts[str(raw)] += 1
        dept_files[str(raw)].add(r.file)
    for raw, n in dept_counts.most_common():
        ev = dept_ev.get(raw, {})
        mapped = map_dept(None if raw == "<blank>" else raw)
        scale_rows.append({"type": "DEPARTMENT", "value": repr(raw), "maps_to": mapped, "count": n,
                           "files": "; ".join(sorted(dept_files[raw]))[:400],
                           "evidence": f"comment keywords sales/service/parts % = {ev.get('sales_kw_%')}/{ev.get('service_kw_%')}/{ev.get('parts_kw_%')} (n={ev.get('n_comments')})" if ev else "",
                           "decision_needed": "YES" if ("suspected" in mapped or mapped == "Unknown" or raw in DEPT_MAP and DEPT_MAP[raw.strip().lower()] in ("Sales",) and raw.lower() in ("used cars", "sales - used")) else ""})
    for b in brand_check:
        scale_rows.append({"type": "BRAND CHECK (dealer mapping)", "value": b["dealer_code"], "maps_to": b["reference"], "count": sum(b["brands_seen"].values()),
                           "files": str(b["brands_seen"]), "evidence": f"unexpected brands: {b['unexpected']}" if b["unexpected"] else "consistent with reference",
                           "decision_needed": "YES" if b["unexpected"] else ""})
    scales_df = pd.DataFrame(scale_rows)

    # Coverage tab
    cov_df = cov.copy()
    cov_df.insert(0, "dealership_name", dealers.set_index("dealer_code").dealership_name)
    cov_df.insert(1, "status", dealers.set_index("dealer_code").status + dealers.set_index("dealer_code").closed_date.fillna("").map(lambda x: f" {x}" if x else ""))
    cov_df["months_with_zero_rows"] = [", ".join(m for m in MONTHS if cov.loc[c, m] == 0) for c in cov.index]
    cov_df["total_rows"] = cov.sum(axis=1)
    cov_df["revenue_months_present"] = [len(rev_months.get(c, set())) for c in cov.index]
    cov_df["revenue_months_after_close"] = [
        ", ".join(sorted(m for m in rev_months.get(c, set()) if isinstance(cd, str) and m > cd[:7]))
        for c, cd in zip(cov.index, dealers.set_index("dealer_code").closed_date)]
    cov_df = cov_df.reset_index()

    # ---------------- Issues log (rule-generated, every issue carries evidence)
    issues = []
    def add(file, sheet, issue, evidence, severity, action):
        issues.append({"id": f"I{len(issues)+1:03d}", "file": file, "sheet": sheet, "issue": issue, "evidence": evidence,
                       "severity": severity, "proposed_action": action, "status": "open"})

    for p in profiles:
        f, s = p["file_name"], p["sheet_name"]
        if p["sheet_kind"].startswith("NON-FEEDBACK"):
            add(f, s, "Tab contains copied public Google reviews, not dealership-collected feedback", f"{p['data_rows']} rows; header '{p['columns'][:80]}'; banner '{p['banner_rows_text'][:80]}'; dates are relative ('a month ago')", "High", "exclude (keep as separate context only)")
        if p["sheet_kind"].startswith("AGGREGATE"):
            add(f, s, "Aggregate summary only - no row-level feedback, no comments", f"{p['data_rows']} summary rows; header '{p['columns'][:90]}'; total rows at {p['total_summary_rows']}", "High", "ask (request raw Q4 2025 data); flag; use counts only for rating-distribution checks")
        if p["sheet_kind"] == "EMPTY sheet":
            add(f, s, "Empty sheet", f"raw_rows={p['raw_rows']}, no values", "Low", "exclude")
        if p["sheet_kind"].startswith("AGGREGATE"):
            present = {map_dept(v.split(":")[0].strip(" '")) for v in p["department_raw_values"].split(";") if v.strip()}
            absent = [d for d in ("Sales", "Service", "Parts") if d not in present]
            if absent:
                add(f, s, f"Department absent from aggregate summary: {', '.join(absent)} (no feedback at all for that department in the period)",
                    f"departments in summary: {p['department_raw_values']}; header '{p['columns'][:90]}'", "High", "ask (was it surveyed?) / log as separate gap")
        if p["sheet_kind"] != "row-level feedback":
            continue  # field-level checks only apply to row-level feedback
        if p["template_change"].startswith("YES"):
            add(f, s, "Template changes partway through the sheet (new header row, different columns)", f"header rows at {p['header_rows']}; templates: {p['columns'][:200]}; {p['rating_scale_observed']}", "High", "fix (parse each block with its own header)")
        if p["total_summary_rows"]:
            add(f, s, "Monthly total / summary rows mixed in with data rows", f"rows {p['total_summary_rows']}", "Medium", "fix (drop summary rows; optionally reconcile counts)")
        if p["banner_row_count"]:
            add(f, s, "Header not on row 1 (banner/title rows above header)", f"header at row {p['header_rows']}; banner: '{p['banner_rows_text'][:150]}'", "Low", "fix (skip banner rows)")
        if p["footer_note_rows"]:
            add(f, s, "Footer note row inside data area", f"row(s) {p['footer_note_rows']}: '{p['footer_note_text'][:150]}'", "Low", "fix (drop note row; keep note as metadata)")
        if p["dates_outside_window"]:
            add(f, s, "Dates outside the Oct 2025 - Sep 2026 window", f"{p['dates_outside_window']} rows: {p['dates_outside_window_rows']}; {p['dates_suspected_year_typo']} of them sit among rows dated exactly one year earlier (suspected year typo)", "High", "flag (suspected typo; confirm before correcting)")
        if p["dates_outside_file_period"] and p["dates_outside_file_period"] != p["dates_outside_window"]:
            add(f, s, "Dates outside the period stated in the file/tab name", f"expected {p['expected_period_from_name']}; {p['dates_outside_file_period']} rows: {p['dates_outside_file_period_rows']}", "High", "flag / ask")
        elif p["dates_outside_file_period"] and p["sheet_kind"] == "row-level feedback" and p["dates_outside_window"] == 0:
            pass
        if p["date_order"].startswith("M/D"):
            add(f, s, "US month/day date order (rest of the folder is day/month)", f"{p['date_part2_gt12']} values have 2nd part > 12 (e.g. {p['date_example_m_d']}), 0 have 1st part > 12; {p['date_ambiguous_dm_count']} more values would silently mis-parse as D/M (e.g. 07/01 -> 7 Jan instead of 1 Jul)", "High", "fix (parse as M/D for this file only)")
        if p["date_order"].startswith("UNPROVEN") and p["date_ambiguous_dm_count"]:
            add(f, s, "Day/month order cannot be proven from values", f"{p['date_ambiguous_dm_count']} ambiguous values, none with a part >12", "Medium", "flag (assume D/M; verify against file period)")
        if p["date_formats"] and ("slash_2digit_year" in p["date_formats"] or "d-Mon-yy" in p["date_formats"] or "d Mon yy" in p["date_formats"] or "excel_serial_number" in p["date_formats"]):
            add(f, s, "Non-standard date representation", p["date_formats"], "Low", "fix (format-specific parser)")
        if p["date_unparsed"]:
            add(f, s, "Unparseable date values", f"{p['date_unparsed']} rows; formats: {p['date_formats']}", "Medium", "flag")
        if p["dates_out_of_sequence"] and p["sheet_kind"] == "row-level feedback" and p["dates_out_of_sequence"] > 0 and p["dates_outside_window"]:
            pass
        if p["inversion_signal"].startswith("SUSPECTED"):
            add(f, s, "Rating scale suspected inverted (1 = best)", f"mean rating of positive-keyword comments {p['inversion_pos_kw_mean']} < negative-keyword comments {p['inversion_neg_kw_mean']} (n pos/neg {p['inversion_n_pos_neg']}); sheet banner: '{p['banner_rows_text'][:120]}'", "High", "fix (reverse to 1=worst) after confirming")
        if "label" in p["rating_kinds"] or "stars" in p["rating_kinds"]:
            add(f, s, "Rating stored as text labels / glyphs, not numbers", f"{p['rating_scale_observed']}; values: {p['rating_distinct_values'][:150]}", "Medium" if "stars" in p["rating_kinds"] else "High", "fix (explicit mapping); ask if label scale has <5 points" if "Bad" in p["rating_distinct_values"] else "fix (explicit mapping)")
        if p["rating_max"] and p["rating_max"] > 5 and p["sheet_kind"] == "row-level feedback":
            add(f, s, "Rating on a 10-point scale", f"{p['rating_scale_observed']}", "High", "fix (rescale to common scale; document method)")
        if p["rating_blank"] and p["sheet_kind"] == "row-level feedback":
            add(f, s, "Missing ratings", f"{p['rating_blank']} of {p['data_rows']} rows", "Medium", "flag (keep row for text analysis, exclude from score metrics)")
        if p["rating_non_parsable"]:
            add(f, s, "Rating values that are not numbers or known labels", f"{p['rating_non_parsable']} rows; distinct: {p['rating_distinct_values'][:150]}", "Medium", "flag")
        if p["test_junk_rows"]:
            add(f, s, "Test / junk rows (e.g. 'TEST', 'asdf', 01/01/1900)", f"rows {p['test_junk_rows']}", "Medium", "exclude")
        if p["exact_duplicate_rows"]:
            add(f, s, "Exact duplicate rows within sheet", f"{p['exact_duplicate_rows']} duplicates: {p['exact_duplicate_row_refs']}", "Medium", "fix (deduplicate) after checking they are not two real visits")
        if p["encoding"] and p["encoding"].startswith("cp1252"):
            add(f, s, "File is Windows-1252 encoded, not UTF-8 (apostrophes break if read as UTF-8)", f"strict UTF-8 decode failed; {p['non_ascii_chars']} non-ASCII chars", "Medium", "fix (read as cp1252)")
        if p["delimiter"] == "semicolon":
            add(f, s, "Semicolon-delimited CSV", f"delimiter detected on header line; field counts {p['field_count_dist']}", "Low", "fix (read with sep=';')")
        if p["field_count_dist"] and p["field_count_dist"].count(":") > 1:
            add(f, s, "CSV rows with inconsistent field counts", f"field count distribution {p['field_count_dist']}", "Medium", "flag")
        if p["comment_mojibake"]:
            add(f, s, "Garbled characters (mojibake / replacement chars) in comments", f"{p['comment_mojibake']} comments", "Low", "fix (decode correctly)")
        if p["unnamed_cols_with_data"]:
            add(f, s, "Columns with data but no header", p["unnamed_cols_with_data"], "Medium", "fix (map by position after template change)")
        pii_txt = p["pii_in_comment_phone"] + p["pii_in_comment_email"] + p["pii_in_comment_rego"] + p["pii_in_comment_name_prefix"]
        if pii_txt:
            add(f, s, "Personal data embedded in free-text comments", f"rows matching phone={p['pii_in_comment_phone']}, email={p['pii_in_comment_email']}, rego/plate={p['pii_in_comment_rego']}, 'Firstname Lastname - ' prefix={p['pii_in_comment_name_prefix']}", "High", "fix (redact before storage and before any LLM call)")
        if p["pii_columns_populated"]:
            add(f, s, "PII columns present", p["pii_columns_populated"], "Medium", "fix (drop name/phone/email/rego; hash customer id)")
        if p["department_source"].startswith("TAB"):
            add(f, s, "No department column; department implied by tab name", f"tab '{s}', {p['data_rows']} rows", "Low", "fix (derive department from tab)")
        if "Unknown" in p["department_mapped"] or "suspected" in p["department_mapped"]:
            add(f, s, "Department values that do not map cleanly", f"raw: {p['department_raw_values'][:150]} -> {p['department_mapped']}", "Medium", "ask / flag")
        if p["months_missing_vs_expected"] and p["sheet_kind"] == "row-level feedback":
            add(f, s, "Months with no rows within the period the file claims to cover", f"expected {p['expected_period_from_name']}; missing {p['months_missing_vs_expected']}", "Medium", "flag / ask")
        if p["comment_staff_paraphrase"]:
            add(f, s, "Comments recorded by staff in third person (paraphrase, not customer verbatim)", f"{p['comment_staff_paraphrase']} comments start with 'Cust/Customer/Said/...'", "Low", "flag (treat sentiment as second-hand)")

    for o in overlaps:
        (fa, sa), (fb, sb) = o["a"], o["b"]
        if o["identical"]:
            add(f"{fa} | {fb}", f"{sa} | {sb}", "Identical sheets (same rows) - one is a copy", f"{o['shared']} of {o['a_rows']}/{o['b_rows']} (date,rating,comment) keys shared; content hash equal", "High", "exclude the copy; ask for the missing period")
        else:
            add(f"{fa} | {fb}", f"{sa} | {sb}", "Overlapping records between two sheets/files of the same dealer", f"{o['shared']} shared (date,rating,comment) keys; sizes {o['a_rows']} vs {o['b_rows']}", "High", "fix (deduplicate; keep latest version) / ask which is authoritative")
    for f, n, ex in cross_tab:
        add(f, "multiple tabs", "Same feedback (date + comment) logged in more than one tab", f"{n} date+comment pairs appear in >1 tab, e.g. {ex}", "Medium", "fix (dedupe; decide department) / ask")
    if cross_dealer_reuse:
        add("(all)", "", "Identical long comments (>=40 chars) appear at more than one dealer", f"{cross_dealer_reuse} distinct comment texts shared across dealers", "Low", "flag (copy-paste or template answers; affects theme counts)")
    # only collisions between DIFFERENT dealers matter here (same-dealer re-sends are covered by the overlap check)
    id_coll_d = ids.groupby("record_id").dealer_code.nunique()
    coll = ids[ids.record_id.isin(id_coll_d[id_coll_d > 1].index)]
    id_coll_files = coll.groupby("file").size().to_dict()
    # pairwise: same ID, different dealers, same non-blank comment
    pairs = coll[coll.ckey.notna()].merge(coll[coll.ckey.notna()], on="record_id")
    pairs = pairs[pairs.dealer_code_x < pairs.dealer_code_y]
    same_content = int((pairs.ckey_x == pairs.ckey_y).sum())
    for f, n in id_coll_files.items():
        add(f, "", "Record IDs reuse the same numbering as another dealer's files (not globally unique)",
            f"{n} rows whose ID also appears in another file; colliding IDs with identical comment text: {same_content} -> shared numbering scheme, not copied rows",
            "Low", "flag (prefix IDs with dealer_code + source file)")
    for f, n in id_dupe_within.items():
        add(f, "", "Duplicate record IDs within file", f"{n} IDs used more than once", "Medium", "flag")
    for b in brand_check:
        if b["unexpected"]:
            add("(dealer mapping)", b["dealer_code"], "Vehicles of brands the dealership is not listed as selling", f"seen {b['brands_seen']} vs reference {b['reference']}", "Medium", "ask (service of other brands is plausible; or wrong dealer mapping)")
    for d in dealers.itertuples():
        if d.status == "Closed":
            after = rl[(rl.dealer_code == d.dealer_code) & rl.date.notna() & (rl.date.astype(str) > str(d.closed_date))]
            rev_after = revenue[(revenue.dealer_code == d.dealer_code) & (revenue.month > str(d.closed_date)[:7])]
            rev_note = f"{len(rev_after)} revenue months after close, all zero: {bool((rev_after.select_dtypes('number') == 0).all().all())}"
            add("; ".join(after.file.unique()) or "reference/dealerships.csv", d.dealer_code, f"Closed site ({d.closed_date}) - feedback dated after closure",
                f"{len(after)} rows dated after closure: {', '.join(f'r{r}={dt}' for r, dt in zip(after.row, after.date))}; {rev_note}; footer says later feedback goes to Bendigo (VIC08)",
                "High", "flag (see year-typo issue; exclude site from post-Mar-2026 trends; check Bendigo for ex-Shepparton customers)")
    for c in cov.index:
        missing = [m for m in MONTHS if cov.loc[c, m] == 0]
        st = dealers.set_index("dealer_code").loc[c]
        if st.status == "Closed":
            missing = [m for m in missing if m <= str(st.closed_date)[:7]]
        if missing:
            add("(coverage)", c, "Dealer has months with zero feedback rows", f"missing: {', '.join(missing)}; total rows {int(cov.loc[c].sum())}", "Medium" if len(missing) < 3 else "High", "ask (request missing data) / flag in trend analysis")
    # department mix per dealer (a missing department skews dealer comparisons)
    rl["dept_std"] = rl.dept_raw.map(map_dept)  # dept_raw already holds the tab name where there is no column
    mix = pd.crosstab(rl.dealer_code, rl.dept_std.str.replace(" (suspected)", "", regex=False), normalize="index")
    for c, r in mix.iterrows():
        low = [f"{d} {100*r.get(d, 0):.1f}%" for d in ("Sales", "Service", "Parts") if r.get(d, 0) < 0.05]
        if low:
            n = int((rl.dealer_code == c).sum())
            add("; ".join(rl[rl.dealer_code == c].file.unique())[:200], c, "Department almost absent from this dealer's feedback (collection gap or bias)",
                f"{n} rows; share: {', '.join(low)}", "Medium", "ask (are Sales/Parts surveyed?) / flag in dealer comparisons")
    rev_cols = {"Sales": "sales_revenue_aud", "Service": "service_revenue_aud", "Parts": "parts_revenue_aud"}
    rev_sum = {d: revenue[c].sum() for d, c in rev_cols.items()}
    add("reference/monthly_revenue.csv", "", "Revenue concentrated in Sales (one-off vehicle purchases), which dominates any revenue-at-risk total",
        "; ".join(f"{d} ${v/1e6:.1f}m ({100*v/sum(rev_sum.values()):.0f}%)" for d, v in rev_sum.items()) + " over Oct 2025 - Sep 2026",
        "Medium", "fix (separate recurring Service + Parts from Sales in Q3 sizing)")
    for f, (code, ev) in alias_files.items():
        n = int(prof_df[prof_df.file_name == f].data_rows.sum())
        add(f, "", "Dealer mapped by alias, not by suburb name (needs decision)",
            f"{n} rows -> {code} {dealers.set_index('dealer_code').loc[code, 'dealership_name']}; {ev}", "Medium", "fix (map by alias) / ask head office to confirm")
    for f in lockfiles:
        add(f, "", "Excel lock/owner file in feedback folder (a workbook is open in Excel)", "file name starts with '~$'; not readable as a workbook", "Low", "exclude (ignore ~$ files in the loader)")
    if errors:
        for f, e in errors:
            add(f, "", "Could not read file/sheet", e, "High", "ask")

    # Volume / collection-bias signal (Noosa-type cards): share of top-box ratings and responses per 100 transactions
    vol = rl.groupby("dealer_code").agg(rows=("row", "size"), mean_rating=("rating_num", "mean"))
    rev_tot = revenue.groupby("dealer_code")[["vehicles_sold", "service_jobs", "parts_transactions"]].sum().sum(axis=1)
    vol["responses_per_1000_txn"] = (1000 * vol.rows / rev_tot).round(2)
    med = vol.responses_per_1000_txn.median()
    for c, r in vol.iterrows():
        if r.responses_per_1000_txn < med / 3:
            sub = prof_df[prof_df.dealer_code.str.startswith(c)]
            add("; ".join(sub.file_name.unique())[:200] or c, c, "Very low response volume relative to transactions (possible selective collection)", f"{int(r.rows)} rows vs median {med:.1f} per 1000 txn (this site {r.responses_per_1000_txn}); mean rating {r.mean_rating:.2f}", "Medium", "flag (low n / possible bias; widen confidence intervals)")

    iss_df = pd.DataFrame(issues)
    sev_rank = {"High": 0, "Medium": 1, "Low": 2}
    # collapse the same issue repeated across tabs of one file into a single row
    iss_df = (iss_df.groupby(["file", "issue"], sort=False)
              .agg(sheet=("sheet", lambda x: ", ".join(dict.fromkeys(map(str, x)))),
                   evidence=("evidence", lambda x: " || ".join(x)[:1500] if len(x) == 1 else " || ".join(f"[{e}]" for e in x)[:1500]),
                   severity=("severity", lambda x: min(x, key=sev_rank.get)),
                   proposed_action=("proposed_action", "first"), status=("status", "first"), n_sheets=("sheet", "size"))
              .reset_index())
    multi = iss_df.n_sheets > 1
    iss_df.loc[multi, "evidence"] = iss_df[multi].apply(lambda r: f"{r.n_sheets} tabs: {r.evidence}", axis=1)
    iss_df.insert(0, "id", "")
    iss_df = iss_df[["id", "file", "sheet", "issue", "evidence", "severity", "proposed_action", "status"]]
    iss_df = iss_df.sort_values(["severity", "file"], key=lambda s: s.map(sev_rank) if s.name == "severity" else s).reset_index(drop=True)
    iss_df["id"] = [f"I{i+1:03d}" for i in range(len(iss_df))]

    dec_df = build_decisions(rl, prof_df, iss_df, dept_ev, scales_df, brand_check)
    linked = {i.strip() for s in dec_df.linked_issue_ids for i in s.split(",") if i.strip()}

    # ---------------- quality checks
    checks = []
    checks.append(("every decision has evidence and an implementation note", bool(dec_df.evidence.str.len().gt(0).all() and dec_df.implementation.str.len().gt(0).all())))
    checks.append(("every decision links >= 1 issue and every selector matched", bool(dec_df.linked_issue_ids.str.len().gt(0).all() and (dec_df.selectors_matched == dec_df.n_selectors).all())))
    checks.append(("all linked issue ids exist in Issues_Log", linked <= set(iss_df.id)))
    filled = dec_df[["evidence", "implementation", "limit_or_risk"]].apply(lambda c: c.str.contains(r"\{\w+\}|MISSING")).any().any()
    checks.append(("no unfilled placeholders / MISSING in decisions", not filled))
    dec_ids = set(dec_df.id)
    checks.append(("every Verification_Plan and Data_Requests row links an existing decision",
                   {v["decision"] for v in VERIFICATION_PLAN + DATA_REQUESTS} <= dec_ids))
    referenced = set(re.findall(r"\b([VR]\d+)\b", " ".join(dec_df.decision)))
    checks.append(("every V#/R# cited in a decision exists", referenced <= {v["id"] for v in VERIFICATION_PLAN + DATA_REQUESTS}))
    feedback_files = {p.name for p in files}
    checks.append(("every feedback file in manifest", feedback_files <= set(man_df.file_name)))
    checks.append(("data files in folder == 52 (pack README)", len(files) - len(lockfiles) == 52))
    checks.append(("manifest rows == file_profile rows + lock files + errors", len(man_df) == len(prof_df) + len(errors) + len(lockfiles)))
    checks.append(("no unreadable data files", not errors))
    checks.append(("manifest row_count total == file_profile data_rows total", int(man_df.row_count.fillna(0).sum()) == int(prof_df.data_rows.sum())))
    checks.append(("all 24 dealers covered by some file", set(dealers.dealer_code) <= set(rl.dealer_code)))
    checks.append(("every issue has evidence", bool(iss_df.evidence.astype(str).str.len().gt(0).all())))
    checks_df = pd.DataFrame(checks, columns=["check", "passed"])

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        man_df.to_excel(xw, sheet_name="Manifest", index=False)
        prof_df.to_excel(xw, sheet_name="File_Profile", index=False)
        pd.DataFrame(coldict).to_excel(xw, sheet_name="Column_Dictionary", index=False)
        scales_df.to_excel(xw, sheet_name="Scales_and_Labels", index=False)
        iss_df.to_excel(xw, sheet_name="Issues_Log", index=False)
        cov_df.to_excel(xw, sheet_name="Dealer_Coverage", index=False)
        dec_df.drop(columns=["n_selectors", "selectors_matched"]).to_excel(xw, sheet_name="Decisions_Log", index=False)
        pd.DataFrame(VERIFICATION_PLAN).to_excel(xw, sheet_name="Verification_Plan", index=False)
        pd.DataFrame(DATA_REQUESTS).to_excel(xw, sheet_name="Data_Requests", index=False)
        checks_df.to_excel(xw, sheet_name="QA_Checks", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for col_cells in ws.columns:
                w = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells[:200])
                ws.column_dimensions[col_cells[0].column_letter].width = min(max(10, w + 2), 60)

    print(f"files={len(files)} sheets={len(prof_df)} data_rows={int(prof_df.data_rows.sum())} issues={len(iss_df)} errors={len(errors)}")
    print(checks_df.to_string(index=False))
    if not checks_df.passed.all():
        sys.exit(1)


if __name__ == "__main__":
    main()
