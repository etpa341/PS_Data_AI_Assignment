"""
make_spot_check.py - 10 silver rows to check by hand against the source files.

One random included row (seed 42) from each case a parsing rule handles specially, plus 5 random included rows
from the other files:
  Geelong 2025          reversed 1-5 scale fixed before normalising (D1)
  Liverpool Q3          US month/day dates
  Ringwood Q2           1-10 score column (D2)
  Shepparton corrected  date moved back a year (D4)
  Bendigo "S"           department inferred as Sales (D3)
Writes outputs/spot_check.xlsx with empty pass_fail (dropdown) and notes columns. Refuses to overwrite an
existing sheet (it may hold checks); pass --force to rebuild. No API calls.   Run:  python make_spot_check.py
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs" / "spot_check.xlsx"
SEED = 42
COLS = ["source_file", "source_sheet", "source_row", "date_raw", "feedback_date", "department", "rating_raw",
        "rating_scale", "score_norm", "comment_clean"]


def main():
    if OUT.exists() and "--force" not in sys.argv:
        sys.exit(f"{OUT} already exists and may hold checks. Pass --force to rebuild it.")
    s = pd.read_csv(ROOT / "data" / "silver" / "feedback_silver.csv", dtype=str, keep_default_na=False)
    inc = s[s["excluded"] == "False"]
    groups = [
        ("Geelong 2025 (reversed scale fixed, D1)", (inc["source_file"] == "Geelong Feedback.xlsx") & (inc["source_sheet"] == "2025")),
        ("Liverpool Q3 (US month/day dates)", inc["source_file"] == "liverpool_q3_2026.csv"),
        ("Ringwood Q2 (1-10 score, D2)", inc["source_file"] == "Ringwood feedback Q2 2026.xlsx"),
        ("Shepparton, date corrected (D4)", inc["source_file"].str.startswith("Shepparton") & (inc["date_corrected"] == "True")),
        ('Bendigo "S" -> Sales, inferred (D3)', (inc["source_file"] == "bendigo feedback.csv") & (inc["dept_inferred"] == "True")),
    ]
    picks = []
    special_files = set()
    for why, mask in groups:
        g = inc[mask]
        assert len(g), why
        special_files |= set(g["source_file"])
        picks.append(g.sample(1, random_state=SEED).assign(why_selected=why))
    other = inc[~inc["source_file"].isin(special_files)]
    picks.append(other.sample(5, random_state=SEED).assign(why_selected="random (other files)"))
    sample = pd.concat(picks).reset_index(drop=True)
    assert len(sample) == 10 and sample["feedback_id"].is_unique

    wb = Workbook()
    ws = wb.active
    ws.title = "Spot_check"
    header = ["check_id", "why_selected", "feedback_id"] + COLS + ["pass_fail", "notes"]
    ws.append(header)
    for i, r in sample.iterrows():
        ws.append([i + 1, r["why_selected"], r["feedback_id"]] + [r[c] for c in COLS] + [None, None])
    dv = DataValidation(type="list", formula1='"pass,fail"', allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = "Pick pass or fail", "Not in list", True
    ws.add_data_validation(dv)
    col = {h: i + 1 for i, h in enumerate(header)}
    letter = lambda h: ws.cell(1, col[h]).column_letter
    dv.add(f"{letter('pass_fail')}2:{letter('pass_fail')}{len(sample) + 1}")
    widths = {"check_id": 8, "why_selected": 30, "feedback_id": 18, "source_file": 34, "source_sheet": 12, "source_row": 10,
              "date_raw": 20, "feedback_date": 13, "department": 11, "rating_raw": 16, "rating_scale": 14, "score_norm": 10,
              "comment_clean": 70, "pass_fail": 10, "notes": 40}
    for h, w in widths.items():
        ws.column_dimensions[letter(h)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}: {len(sample)} rows")
    print(sample[["why_selected", "source_file", "source_sheet", "source_row", "date_raw", "feedback_date", "department",
                  "rating_raw", "rating_scale", "score_norm"]].to_string(index=False))


if __name__ == "__main__":
    main()
