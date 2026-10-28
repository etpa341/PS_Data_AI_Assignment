"""
make_labelling_sheet.py - hand-labelling sample for the AI evaluation (SPEC section 10).

- 120 comments, stratified random across dealer x department (proportional allocation,
  largest remainder), seed 42, customer-voice rows with text only, included rows only, shuffled.
- No rating, score or AI columns, so labels are made blind.
- Tab Voice_check: 30 Essendon / Ipswich rows (15 tagged staff_paraphrased, 15 not) to test the voice tag.
- Refuses to overwrite an existing sheet (it holds hand labels); pass --force to rebuild.

No LLM calls.  Run after 02_silver.py:  python make_labelling_sheet.py
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
SILVER = ROOT / "data" / "silver" / "feedback_silver.csv"
OUT = ROOT / "labelling" / "labelling_sheet.xlsx"
SEED, N, N_VOICE = 42, 120, 15

THEMES = ["Staff attitude", "Communication and updates", "Wait time and turnaround", "Price and billing",
          "Work quality and fix", "Vehicle or product", "Process and facilities", "Other"]
SENTIMENTS = ["positive", "neutral", "negative", "mixed"]
# Short definitions drafted for this sheet (SPEC section 9 names the themes but does not define them)
THEME_DEFS = {
    "Staff attitude": "How staff treated the customer: friendly, rude, helpful, dismissive, pushy, respectful.",
    "Communication and updates": "Being kept informed: call-backs, progress updates, explanations of work or paperwork, messages passed on.",
    "Wait time and turnaround": "Time taken: waiting on the day, how long the job or delivery took, delays, parts on back-order.",
    "Price and billing": "Cost and invoices: price level, value, quotes vs bill, add-ons, fees, upselling of extras.",
    "Work quality and fix": "Whether the job was done right: fixed first time, repeat faults, damage, cleanliness of the returned car.",
    "Vehicle or product": "The car or part itself: model features, reliability, product faults not caused by the dealer's work.",
    "Process and facilities": "Booking and handover processes, paperwork, waiting area, loan cars, parking, Wi-Fi, opening hours.",
    "Other": "None of the above, or too vague to assign (e.g. 'ok', 'good') - give sentiment only.",
}
RULES = [  # SPEC.md section 9, verbatim
    "Themes (one primary, optional second): Staff attitude; Communication and updates; Wait time and turnaround; Price and billing; Work quality and fix; Vehicle or product; Process and facilities; Other.",
    "Sentiment: positive, neutral, negative, mixed.",
    "Blank comments are tagged no_comment and never sent.",
    "Vague comments like \"ok\" go to Other with sentiment only.",
    "If two themes compete, the primary is the one that best explains the rating.",
]


def allocate(sizes, n):
    """Proportional allocation with largest remainder; returns {stratum: k}."""
    quota = sizes / sizes.sum() * n
    k = quota.astype(int)
    rest = n - k.sum()
    order = (quota - k).sort_values(ascending=False, kind="stable").index[:rest]
    k[order] += 1
    return k


def style(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    if OUT.exists() and "--force" not in sys.argv:
        sys.exit(f"{OUT} already exists and may hold hand labels. Pass --force to rebuild it.")
    s = pd.read_csv(SILVER, dtype=str, keep_default_na=False)
    pool = s[(s["excluded"] == "False") & (s["has_text"] == "True") & (s["voice"] == "customer")]
    sizes = pool.groupby(["dealer_code", "department"]).size()
    k = allocate(sizes, N)
    parts = [pool[(pool["dealer_code"] == d) & (pool["department"] == dp)].sample(int(n), random_state=SEED)
             for (d, dp), n in k.items() if n > 0]
    sample = pd.concat(parts).sample(frac=1, random_state=SEED).reset_index(drop=True)
    assert len(sample) == N and sample["feedback_id"].is_unique

    ei = s[(s["excluded"] == "False") & (s["has_text"] == "True")
           & s["source_file"].isin(["ESSENDON - feedback register.xlsx", "IPS_feedback_log.xlsx"])]
    staff = ei[ei["voice"] == "staff_paraphrased"].sample(N_VOICE, random_state=SEED)
    cust = ei[(ei["voice"] == "customer") & ~ei["feedback_id"].isin(sample["feedback_id"])].sample(N_VOICE, random_state=SEED)
    voice = pd.concat([staff, cust]).sample(frac=1, random_state=SEED).reset_index(drop=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Labels"
    ws.append(["feedback_id", "dealer_code", "department", "comment_clean",
               "my_primary_theme", "my_secondary_theme", "my_sentiment", "notes"])
    for r in sample.itertuples():
        ws.append([r.feedback_id, r.dealer_code, r.department, r.comment_clean, None, None, None, None])
    last = N + 1
    dv_theme = DataValidation(type="list", formula1='"' + ",".join(THEMES) + '"', allow_blank=True)
    dv_sent = DataValidation(type="list", formula1='"' + ",".join(SENTIMENTS) + '"', allow_blank=True)
    for dv in (dv_theme, dv_sent):
        dv.error, dv.errorTitle, dv.showErrorMessage = "Pick a value from the list", "Not in list", True
        ws.add_data_validation(dv)
    dv_theme.add(f"E2:F{last}")
    dv_sent.add(f"G2:G{last}")
    style(ws, {"A": 18, "B": 11, "C": 11, "D": 80, "E": 26, "F": 26, "G": 13, "H": 30})

    wd = wb.create_sheet("Definitions")
    wd.append(["item", "text"])
    for t in THEMES:
        wd.append([f"Theme: {t}", THEME_DEFS[t]])
    wd.append(["Sentiment values", ", ".join(SENTIMENTS)])
    for i, rule in enumerate(RULES, 1):
        wd.append([f"SPEC 9 rule {i}", rule])
    wd.append(["Note", "Theme definitions are short working definitions drafted for this sheet; the rules are copied from SPEC.md section 9. "
                       "Label from the comment text only; the rating is deliberately not shown. [STAFF], [NAME], [PHONE], [EMAIL], [REGO] are redactions."])
    style(wd, {"A": 32, "B": 110})

    wv = wb.create_sheet("Voice_check")
    wv.append(["feedback_id", "dealer_code", "comment_clean", "my_voice"])
    for r in voice.itertuples():
        wv.append([r.feedback_id, r.dealer_code, r.comment_clean, None])
    dv_voice = DataValidation(type="list", formula1='"customer,staff"', allow_blank=True)
    wv.add_data_validation(dv_voice)
    dv_voice.add(f"D2:D{len(voice) + 1}")
    style(wv, {"A": 18, "B": 11, "C": 90, "D": 13})

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    # summary (counts only)
    print(f"Labels: {len(sample)} rows from a pool of {len(pool)} customer-voice comments; "
          f"strata {int((k > 0).sum())} of {len(k)} sampled; dealers {sample['dealer_code'].nunique()}")
    print("by department:", sample["department"].value_counts().to_dict(), "| pool share:",
          (pool["department"].value_counts(normalize=True) * 100).round(1).to_dict())
    print("strata with 0 rows allocated:", int((k == 0).sum()), "| largest allocations:", k.sort_values(ascending=False).head(3).to_dict())
    print(f"Voice_check: {len(voice)} rows ({len(staff)} tagged staff, {len(cust)} untagged), "
          f"files {voice['feedback_id'].map(lambda i: i.split('-')[0]).value_counts().to_dict()}")


if __name__ == "__main__":
    main()
